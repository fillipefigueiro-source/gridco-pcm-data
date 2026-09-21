# -*- coding: utf-8 -*-
"""Inicio Semana 39 — Desafios (21 a 25/09/2026), padrao visual S36. R01.
Copia de _s39_inicio.py com as melhorias aprovadas pelos revisores.
Dados: _semana39_dados.json + _s39_mpas_obs.json (API Fracttal em 21/09)."""
import io, json
import re as _re
from collections import defaultdict
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(0x19, 0x15, 0x28)
LIME = RGBColor(0xA9, 0xDB, 0x21)
BODY = RGBColor(0x50, 0x4C, 0x63)
INK = RGBColor(0x23, 0x20, 0x3A)
CALL = RGBColor(0xF0, 0xF1, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
AMBER = RGBColor(0xB9, 0x77, 0x0E)
RED = RGBColor(0xB3, 0x26, 0x1E)
F = 'Calibri'
RODAPE = 'Semana 39 - 21/09 a 25/09 - PCM Grid Co.'

D = json.load(io.open('_semana39_dados.json', encoding='utf-8'))
G = D['geral']
CLIS = D['clientes']
NOME = {'Utragaz': 'Ultragaz'}  # typo no cadastro do Fracttal

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def tx(s, x, y, w, h, texto, sz, cor=INK, bold=False, alinha=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    linhas = texto if isinstance(texto, list) else [texto]
    for i, ln in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinha
        if i:
            p.space_before = Pt(4)
        partes = ln if isinstance(ln, list) else [(ln, cor, bold, sz)]
        for t_, c_, b_, z_ in partes:
            r = p.add_run(); r.text = t_
            r.font.name, r.font.size, r.font.bold = F, Pt(z_), b_
            r.font.color.rgb = c_
    return tb


def caixa(s, x, y, w, h, fill, radius=0.06):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                            Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background(); sh.shadow.inherit = False
    sh.adjustments[0] = radius
    return sh


def retang(s, x, y, w, h, fill):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                            Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh


def slide(titulo, kicker):
    s = prs.slides.add_slide(BLANK)
    tx(s, 0.55, 0.32, 12.2, 0.3, kicker.upper(), 12, BODY, True)
    tx(s, 0.55, 0.62, 12.2, 0.7, titulo, 28, NAVY, True)
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.28),
                            Inches(12.23), Emu(9525))
    ln.fill.solid(); ln.fill.fore_color.rgb = LIME
    ln.line.fill.background(); ln.shadow.inherit = False
    tx(s, 0.55, 7.08, 8.0, 0.3, RODAPE, 9, BODY)
    return s


def kpi(s, x, y, w, valor, rotulo, cor=NAVY, sz=30):
    caixa(s, x, y, w, 1.05, CALL)
    tx(s, x + 0.15, y + 0.1, w - 0.3, 0.55, valor, sz, cor, True)
    tx(s, x + 0.15, y + 0.66, w - 0.3, 0.38, rotulo, 10, BODY)


# V7/V12: cel aceita lista de runs [(texto, cor, bold)] e colunas centralizadas
def tabela(s, x, y, w, cols, largs, linhas, hrow=0.3, sz=10, centro=()):
    t = s.shapes.add_table(len(linhas) + 1, len(cols), Inches(x), Inches(y),
                           Inches(w), Inches(hrow) * (len(linhas) + 1)).table
    t.first_row = False
    t.horz_banding = False
    for i, lg in enumerate(largs):
        t.columns[i].width = Inches(lg)
    for r_ in t.rows:
        r_.height = Inches(hrow)

    def cel(c, runs, fundo, alinha=PP_ALIGN.LEFT, z=sz):
        c.margin_left = c.margin_right = Emu(50000)
        c.margin_top = c.margin_bottom = Emu(12000)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        if fundo is None:
            c.fill.background()
        else:
            c.fill.solid(); c.fill.fore_color.rgb = fundo
        p = c.text_frame.paragraphs[0]
        p.alignment = alinha
        for t_, cor, bold in runs:
            r = p.add_run(); r.text = str(t_)
            r.font.name, r.font.size, r.font.bold = F, Pt(z), bold
            r.font.color.rgb = cor

    for j, h in enumerate(cols):
        cel(t.cell(0, j), [(h, WHITE, True)], NAVY,
            PP_ALIGN.CENTER if j in centro else PP_ALIGN.LEFT)
    for i, linha in enumerate(linhas, 1):
        zebra = CALL if i % 2 == 0 else None
        for j, v in enumerate(linha):
            if isinstance(v, list):
                runs = v
            elif isinstance(v, tuple):
                runs = [v]
            else:
                runs = [(v, BODY, False)]
            cel(t.cell(i, j), runs, zebra,
                PP_ALIGN.CENTER if j in centro else PP_ALIGN.LEFT)
    return t


def usina_curta(u):
    m = _re.match(r'^[^-]+?\s*-\s*(.+?)\s*-\s*[A-Za-z]{2}$', ' '.join(str(u).split()))
    return m.group(1).strip() if m else u


# C9: truncamento em fim de palavra com reticências
def trunc(s_, n):
    s_ = str(s_).strip()
    if len(s_) <= n:
        return s_
    corte = s_[:n].rsplit(' ', 1)[0].rstrip(' ,;·—-')
    return corte + '…'


def nome_cli(c):
    return NOME.get(c, c)


def cli_por(nome):
    return next((c for c in CLIS if c['cli'] == nome), None)


tho = cli_por('Thopen')
ath = cli_por('Athon')
axi = cli_por('Axis')

# ── 1 CAPA ───────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY
bg.line.fill.background(); bg.shadow.inherit = False
retang(s, 0.9, 1.9, 0.6, 0.06, LIME)  # V10
tx(s, 0.9, 2.1, 11.5, 0.4, 'PCM - REUNIÃO DE SEMANA', 15, LIME, True)  # V10
tx(s, 0.9, 2.55, 11.5, 1.0, 'Semana 39 — nossos desafios', 44, WHITE, True)
tx(s, 0.9, 3.75, 11.5, 0.4, '21 a 25 de setembro de 2026', 16,
   RGBColor(0xCF, 0xCB, 0xDA))
tx(s, 0.9, 6.6, 11.5, 0.4,
   'Equipe de PCM - Grid Co. - dados da API Fracttal em 21/09 (R01)', 11,
   RGBColor(0x9B, 0x96, 0xAD))

# ── 2 AGENDA ─────────────────────────────────────────────────────────
s = slide('A semana em uma página', 'PCM - Agenda')
itens = [
    ('1', 'MPM de setembro — última semana', '%d%% (%d de %d tarefas) — o mês fecha na quarta 30/09 (slide 7)'
     % (G['mpm_pct'], G['mpm_fei'], G['mpm_tot'])),
    ('2', 'MPA vencidas', 'subiram para %d OS — Petrolina 2 (Axis) chegou a 104 dias (slide 8)'
     % G['mpa_venc']),
    ('3', 'MPS vencidas', '%d OS, todas Thopen — a OS 6340 (118d) FECHOU; a mais antiga agora é Araçoiaba da Serra 1 (81d) (slide 8)'
     % G['mps_venc']),  # V13
    ('4', 'Handover', '%d OS abertas, todas Thopen — Inhapi 1 fechou; a mais antiga é Coração 1 e 2 (65d) (slide 9)' % G['ho_ab']),
    ('5', 'Corretivas', 'caíram para %d abertas (eram 126); %d com +30 dias — backlog por supervisor (slide 10)'
     % (G['corr_ab'], G['corr_30d'])),
    ('6', 'Pontos por cliente', 'Ultragaz zerou corretivas e fechou MPM — risco concentra em Thopen, Athon, Axis e RenoGrid (slides 4 a 6)'),
    ('7', 'Análise de óleo', '%d OS abertas — manter a logística enviar / rota / usina / laboratório' % G['oleo_ab']),
    ('8', 'Ritmo da semana passada', 'S38 fechou %d OS; entraram %d corretivas novas'
     % (G['fin_s37'], G['corr_criadas_s37'])),
]
y = 1.55
for num, t_, d_ in itens:
    caixa(s, 0.55, y, 0.42, 0.42, LIME, 0.5)
    tx(s, 0.55, y + 0.075, 0.42, 0.3, num, 14, NAVY, True, PP_ALIGN.CENTER)  # V14
    tx(s, 1.15, y + 0.02, 11.4, 0.4,
       [[(t_ + ' — ', INK, True, 13), (d_, BODY, False, 12)]], 13)
    y += 0.66

# ── 3 PANORAMA POR CLIENTE ───────────────────────────────────────────
s = slide('Panorama por cliente — onde está o risco', 'PCM - Portfólio')
# C7: deltas vs S38 | V3: rótulo do 4º KPI
kpi(s, 0.55, 1.5, 2.9, str(G['corr_ab']), 'corretivas abertas (eram 126)', AMBER)
kpi(s, 3.6, 1.5, 2.9, str(G['mpa_venc'] + G['mps_venc']), 'MPA/MPS vencidas (eram 31)', RED)
kpi(s, 6.65, 1.5, 2.9, '%d%%' % G['mpm_pct'], 'MPM de setembro (era 37%)', AMBER)
kpi(s, 9.7, 1.5, 3.05, str(G['backlog_os']),
    'OS paradas há +30 dias, todos os tipos (eram 70)', RED)


def fmt_mpm(c):
    if c['mpm_pct'] is None:
        return ('—', BODY, False)
    # V9: RED <40, AMBER 40–79, GREEN >=80
    cor = GREEN if c['mpm_pct'] >= 80 else (AMBER if c['mpm_pct'] >= 40 else RED)
    return ('%d%%' % c['mpm_pct'], cor, True)


linhas = []
for c in CLIS:
    ant = c['antiga']
    if ant:
        # V7: só o "· NNd" fica vermelho/bold quando >60d
        ant_runs = [('OS %s · %s ' % (ant['os'], usina_curta(ant['u'])[:18]), BODY, False),
                    ('· %dd' % ant['dias'],
                     RED if ant['dias'] > 60 else BODY, ant['dias'] > 60)]
    else:
        ant_runs = '—'
    linhas.append((
        (nome_cli(c['cli']), NAVY, True),
        '%d em %d usinas' % (c['corr_ab'], c['usinas_corr']) if c['corr_ab'] else '—',
        (str(c['corr_30d']), RED if c['corr_30d'] >= 5 else BODY, c['corr_30d'] >= 5),
        '%d / %d' % (c['mpa_ab'], c['mpa_venc']) if c['mpa_ab'] else '—',
        '%d / %d' % (c['mps_ab'], c['mps_venc']) if c['mps_ab'] else '—',
        fmt_mpm(c),
        ant_runs))
tabela(s, 0.55, 2.85, 12.2,
       ['Cliente', 'Corretivas abertas', '+30d', 'MPA ab/venc', 'MPS ab/venc',
        'MPM set', 'OS aberta mais antiga'],
       [1.55, 1.95, 0.7, 1.35, 1.35, 1.0, 4.3], linhas, hrow=0.29,
       centro={2, 3, 4})  # V12

# ── 4 THOPEN ─────────────────────────────────────────────────────────
c = tho
s = slide('Thopen — maior volume do portfólio', 'PCM - Pontos por cliente 1/3')
kpi(s, 0.55, 1.5, 2.9, str(c['corr_ab']), 'corretivas abertas em %d usinas' % c['usinas_corr'], AMBER)
kpi(s, 3.6, 1.5, 2.9, str(c['mpa_venc']), 'MPA vencidas (%d abertas)' % c['mpa_ab'], RED)
kpi(s, 6.65, 1.5, 2.9, str(c['mps_venc']), 'MPS vencidas (%d abertas)' % c['mps_ab'], RED)
kpi(s, 9.7, 1.5, 3.05, '%d%%' % c['mpm_pct'], 'MPM set (%d/%d tarefas)' % (c['mpm_fei'], c['mpm_tot']), AMBER)
tx(s, 0.55, 2.85, 12.2, 0.3, 'Usinas com mais corretivas abertas:', 12, INK, True)
linhas = [((trunc(usina_curta(u['u']), 34), NAVY, True), ('%d OS' % u['n'], AMBER, True))
          for u in c['top_usinas']]
# V4: hrow 0.45 para igualar a base com o box da direita
tabela(s, 0.55, 3.25, 5.6, ['Usina', 'Corretivas'], [4.0, 1.6], linhas,
       hrow=0.45, centro={1})  # V12
caixa(s, 6.55, 3.25, 6.2, 1.85, CALL)
# C1: vitória da semana + nova mais antiga (Thopen e portfólio)
tx(s, 6.8, 3.45, 5.7, 1.6, [
    [('Vitória da semana: ', NAVY, True, 12),
     ('a OS 6340 (MPS de Aparecida do Taboado, aberta desde 19/05) foi FECHADA '
      'após ~4 meses.', GREEN, True, 12)],
    [('Mais antiga da Thopen agora: ', NAVY, True, 11),
     ('OS %s — %s (MPA Caixa d\'água), ' % (c['antiga']['os'],
      usina_curta(c['antiga']['u'])), BODY, False, 11),
     ('%d dias' % c['antiga']['dias'], RED, True, 11),
     ('. No portfólio, a mais antiga é a ', BODY, False, 11),
     ('OS %s — %s (Axis), %d dias' % (axi['antiga']['os'],
      usina_curta(axi['antiga']['u']), axi['antiga']['dias']), RED, True, 11),
     ('.', BODY, False, 11)]], 12)
# V4: box do desafio sobe para y 5.25 com altura 1.35
caixa(s, 0.55, 5.25, 12.2, 1.35, CALL)
lider = ' e '.join(usina_curta(u['u']) for u in c['top_usinas'][:2])
tx(s, 0.75, 5.42, 11.8, 1.05, [
    [('Desafio da semana: ', NAVY, True, 12),
     ('fechar o mês — MPM em %d%% com %d usinas pendentes; '
      % (c['mpm_pct'], len(c['mpm_pend_usinas'])), INK, False, 12),
     ('%d MPA e %d MPS vencidas (Araçoiaba SPDA e Mandaguaçu caixa d\'água com 81d) '
      'e %d handovers abertos (slide 9).'
      % (c['mpa_venc'], c['mps_venc'], c['ho_ab']), INK, False, 12)],
    [('%s seguem liderando em corretivas — revisar na programação de hoje.' % lider,
      BODY, False, 11)]], 12)

# ── 5 ATHON + AXIS ───────────────────────────────────────────────────
s = slide('Athon e Axis — atenção concentrada', 'PCM - Pontos por cliente 2/3')


def bloco_cli(x, c, bullets, desafio):
    caixa(s, x, 1.55, 5.95, 5.3, CALL)
    tx(s, x + 0.25, 1.8, 5.4, 0.4, nome_cli(c['cli']), 18, NAVY, True)
    kpis = [('%d' % c['corr_ab'], 'corretivas abertas', AMBER),
            ('%d' % c['corr_30d'], 'com +30 dias', RED if c['corr_30d'] else NAVY),
            ('%d%%' % (c['mpm_pct'] or 0), 'MPM setembro',
             RED if (c['mpm_pct'] or 0) < 40 else AMBER)]
    kx = x + 0.25
    for v, r_, cor in kpis:
        caixa(s, kx, 2.35, 1.75, 0.95, WHITE)
        tx(s, kx + 0.12, 2.44, 1.5, 0.5, v, 24, cor, True)
        tx(s, kx + 0.12, 2.95, 1.55, 0.3, r_, 9, BODY)
        kx += 1.88
    # V6: recuo pendente — marcador e texto em textboxes separados
    by = 3.5
    for b in bullets:
        tx(s, x + 0.25, by, 0.2, 0.2, '•', 11, NAVY, True)
        tx(s, x + 0.5, by, 5.15, 0.6, b, 11, INK)
        nl = -(-len(b) // 78)
        by += nl * 0.19 + 0.10
    # V1: mini-box "Desafio da semana" na base do cartão
    caixa(s, x + 0.25, 5.6, 5.45, 1.05, WHITE)
    tx(s, x + 0.42, 5.72, 5.12, 0.85,
       [[('Desafio da semana: ', NAVY, True, 11), (desafio, INK, False, 11)]], 11)


bloco_cli(0.55, ath, [
    'OS 8338 — Timon 1: trackers parados há %d dias (desde 24/06).' % ath['antiga']['dias'],
    # C10: dois fatos separados
    'Marabá 2 e Santa Maria do Pará lideram em corretivas abertas (5 cada).',
    '%d das %d corretivas já passaram de 30 dias.' % (ath['corr_30d'], ath['corr_ab']),
    'MPM de setembro em %d%% — %d usinas pendentes até 30/09.'
    % (ath['mpm_pct'], len(ath['mpm_pend_usinas'])),
    'MPA de Mãe do Rio 1 venceu 17/09 — em execução (plano de ação com a Athon).',
], 'fechar as %d usinas de MPM e decidir as %d corretivas +30d.'
   % (len(ath['mpm_pend_usinas']), ath['corr_30d']))
bloco_cli(6.8, axi, [
    'As %d MPA abertas estão TODAS vencidas — inclui Marialva 1 (crítica, pendente de envio de relatório).'
    % axi['mpa_ab'],
    # C2: desligamento remarcado para 10/09
    'Petrolina 2 chegou a %d dias (desde 08/06) — o desligamento foi remarcado para '
    '10/09 (Caio/Axis): confirmar se ocorreu e se os ensaios foram concluídos.'
    % axi['mpa_antiga']['dias'],
    'MPM de setembro em %d%% — %d usinas pendentes.'
    % (axi['mpm_pct'], len(axi['mpm_pend_usinas'])),
    '%d corretivas abertas em %d usinas — Marialva 1 concentra 3.'
    % (axi['corr_ab'], axi['usinas_corr']),
], 'confirmar se o desligamento de Petrolina 2 (remarcado para 10/09) foi '
   'executado e concluir os ensaios.')

# ── 6 DEMAIS CLIENTES ────────────────────────────────────────────────
s = slide('Demais clientes — um ponto por linha', 'PCM - Pontos por cliente 3/3')
resto = [
    ('RenoGrid', '5 MPA abertas, TODAS vencidas e Muito Crítico na Gerencial (cluster com 1 técnico só); MPM em 0%.', RED),
    ('GD Energy', 'MPM em 20%; exaustores do trafo de Guajirú seguem abertos há 62 dias (2 corretivas +30d).', RED),
    ('Alves Lima', 'MPM de setembro em 0% (12 tarefas em Morada Nova 1) — corretivas zeradas.', RED),
    ('Sal Energia', 'MPM subiu de 8% para 44% — bom avanço; reparo de estruturas de Aquiraz 2 chegou a 32d.', AMBER),
    ('Greenyellow', 'MPM em 94%; seguem as 3 MPA vencidas — Cedro 1 (79d) com falta de ferramentas relatada.', AMBER),
    ('2C', 'MPM em 68% (entraram usinas novas no plano); chamado STI de Sete Lagoas 2 há 55d; MPA de Tupi Paulista vencida.', AMBER),
    ('Ultragaz', 'ZEROU as corretivas e fechou a MPM em 100% — restam só as 2 MPA vencidas de Ibirapuã (41d).', GREEN),
    ('Semp', 'MPM 100%; análise do descasamento de geração de Tucano 1 completou 30 dias.', GREEN),
]
y = 1.65
for nome, txt_, cor in resto:
    caixa(s, 0.55, y, 12.2, 0.52, CALL, 0.12)
    # V11: barra insetada para não vazar no canto arredondado
    retang(s, 0.55, y + 0.06, 0.09, 0.40, cor)
    tx(s, 0.85, y + 0.11, 1.85, 0.3, nome, 12, NAVY, True)
    tx(s, 2.75, y + 0.12, 9.9, 0.3, txt_, 11, INK)
    y += 0.62

# ── 7 MPM SETEMBRO ───────────────────────────────────────────────────
s = slide('MPM de setembro — última semana, %d%% feito' % G['mpm_pct'],
          'PCM - Preventiva mensal')
mpm_falta = G['mpm_tot'] - G['mpm_fei']
cem = [nome_cli(c['cli']) for c in CLIS if c['mpm_tot'] and c['mpm_pct'] == 100]
baixo = [nome_cli(c['cli']) for c in CLIS if c['mpm_tot'] and c['mpm_pct'] <= 20]
kpi(s, 0.55, 1.5, 2.9, '%d%%' % G['mpm_pct'], '%d de %d tarefas' % (G['mpm_fei'], G['mpm_tot']), AMBER)
# V8: KPIs 2 e 3 como número grande; KPI 4 com valor 645
kpi(s, 3.6, 1.5, 2.9, str(len(cem)), 'clientes já em 100%: Ultragaz · Semp', GREEN)
kpi(s, 6.65, 1.5, 2.9, str(len(baixo)), 'em 0–20%: RenoGrid · GD · Alves Lima', RED)
kpi(s, 9.7, 1.5, 3.05, str(mpm_falta),
    'tarefas p/ fechar 100% até 30/09 (~81/dia útil)', NAVY)
mpms = sorted([c for c in CLIS if c['mpm_tot']], key=lambda c: c['mpm_pct'])
linhas = []
for c in mpms:
    base = ', '.join(usina_curta(u) for u in c['mpm_pend_usinas'][:4])
    onde_falta = trunc(base, 68)  # C9: fim de palavra
    if len(c['mpm_pend_usinas']) > 4:
        onde_falta += ' +%d' % (len(c['mpm_pend_usinas']) - 4)
    linhas.append(((nome_cli(c['cli']), NAVY, True),
                   '%d/%d' % (c['mpm_fei'], c['mpm_tot']),
                   fmt_mpm(c),
                   str(len(c['mpm_pend_usinas'])),
                   onde_falta))
tabela(s, 0.55, 2.85, 12.2,
       ['Cliente', 'Feitas', '%', 'Usinas pend.', 'Onde falta'],
       [1.7, 1.1, 0.8, 1.2, 7.4], linhas, hrow=0.31,
       centro={1, 2, 3})  # V12
# C5: teste de ritmo sob a tabela
tx(s, 0.55, 2.85 + 0.31 * (len(linhas) + 1) + 0.12, 12.2, 0.3,
   [[('Teste de ritmo: ', NAVY, True, 10),
     ('%d tarefas em 8 dias úteis ≈ 81/dia, o dobro da semana passada — '
      'decidir hoje o que fecha e o que será justificado.' % mpm_falta,
      BODY, False, 10)]], 10)

# ── 8 MPA E MPS POR CLIENTE ──────────────────────────────────────────
OBS = json.load(io.open('_s39_mpas_obs.json', encoding='utf-8'))
CRIT_RANK = {'muito crítico': 0, 'crítico': 1, 'alto': 2, 'médio': 3, 'baixo': 4}
CRIT_COR = {0: RED, 1: RED, 2: AMBER, 3: AMBER, 4: GREEN}


def data_max(txt):
    """C8: data mais recente dd/mm/aaaa dentro do texto; sem data = mais antiga."""
    ds = _re.findall(r'(\d{2})/(\d{2})/(\d{4})', txt or '')
    if not ds:
        return (0, 0, 0)
    return max((int(a), int(m), int(d)) for d, m, a in ds)


def crit_obs(cli):
    ent = OBS.get(cli, [])
    if not ent:
        return None
    rank = min(CRIT_RANK.get(e['crit'].lower(), 9) for e in ent)
    top = [e for e in ent if CRIT_RANK.get(e['crit'].lower(), 9) == rank]
    # C8: entre as usinas de criticidade máxima, a observação com data mais recente
    e = max(top, key=lambda e_: data_max(e_['obs']))
    partes = [p.strip() for p in e['obs'].replace('\n', ' ').split('•') if p.strip()]
    obs = partes[-1] if partes else ''
    if obs:
        obs = '%s: %s' % (usina_curta(e['u']), obs)
    return (e['crit'] or '—', CRIT_COR.get(rank, BODY), rank <= 1, obs)


def fmt_dias(d):
    if d > 0:
        return ('%dd' % d, RED if d > 45 else AMBER, d > 45)
    return ('no prazo', GREEN, False) if d < 0 else ('hoje', AMBER, False)


def fmt_ant(a):
    if not a:
        return '—'
    idade, cor, b = fmt_dias(a['dias'])
    return ('OS %s · %s · %s' % (a['os'], usina_curta(a['u'])[:16], idade), cor, b)


s = slide('MPA e MPS — abertas e vencidas por cliente', 'PCM - Anuais e semestrais')
kpi(s, 0.55, 1.5, 2.9, str(G['mpa_ab']), 'OS de MPA abertas', AMBER)
kpi(s, 3.6, 1.5, 2.9, str(G['mpa_venc']), 'MPA já vencidas', RED)
kpi(s, 6.65, 1.5, 2.9, str(G['mps_ab']), 'OS de MPS abertas', AMBER)
kpi(s, 9.7, 1.5, 3.05, str(G['mps_venc']), 'MPS vencidas (todas Thopen)', RED)
# V2: nota de MPS sob os KPIs (colunas MPS saem da tabela — MPS é só Thopen)
tx(s, 0.55, 2.62, 12.2, 0.25,
   [[('MPS: ', NAVY, True, 10),
     ('%d abertas / %d vencidas — todas Thopen; mais antiga OS %s · %s · %dd'
      % (G['mps_ab'], G['mps_venc'], tho['mps_antiga']['os'],
         usina_curta(tho['mps_antiga']['u']), tho['mps_antiga']['dias']),
      BODY, False, 10)]], 10)
camp = [c for c in CLIS if c['mpa_ab'] or c['mps_ab']]
camp.sort(key=lambda c: -(c['mpa_venc'] + c['mps_venc']))
linhas = []
for c in camp:
    onde = sorted({usina_curta(u) for u in c['mpa_usinas'] + c['mps_usinas']})
    co = crit_obs(c['cli'])
    base = ', '.join(onde[:2])
    suf = ' +%d' % (len(onde) - 2) if len(onde) > 2 else ''
    onde_txt = trunc(base, 40 - len(suf)) + suf  # C9: "+N" nunca cortado
    linhas.append((
        (nome_cli(c['cli']), NAVY, True),
        ('%d / %d' % (c['mpa_ab'], c['mpa_venc']),
         RED if c['mpa_venc'] >= 3 else BODY, c['mpa_venc'] >= 3),
        fmt_ant(c['mpa_antiga']),
        onde_txt,
        (co[0], co[1], co[2]) if co else '—',
        (trunc(co[3], 95) if co and co[3] else '—')))
# V2: sem colunas de MPS; Observações ~4.6"; fonte 10pt; y 2.85
tabela(s, 0.55, 2.85, 12.2,
       ['Cliente', 'MPA ab/venc', 'MPA mais antiga', 'Onde', 'Criticidade',
        'Observações'],
       [1.3, 1.05, 2.2, 1.9, 1.15, 4.6], linhas, hrow=0.46, sz=10,
       centro={1})  # V12
tx(s, 4.7, 7.08, 8.05, 0.3,
   'Criticidade e observações: Gerencial - PCM_2026_R00, aba MPAS (usina mais '
   'crítica de cada cliente) · 21/09.', 8.5, BODY, False, PP_ALIGN.RIGHT)

# ── 9 HANDOVER ───────────────────────────────────────────────────────
s = slide('Handover — pendências de entrada de usina', 'PCM - Handover')
HO = D['handover']
ho_venc = [h for h in HO if h['antiga'] and h['antiga']['dias'] > 0]
kpi(s, 0.55, 1.5, 2.9, str(G['ho_ab']), 'OS de handover abertas (todas Thopen)', AMBER)
kpi(s, 3.6, 1.5, 2.9, str(len(HO)), 'usinas com handover aberto', NAVY)
kpi(s, 6.65, 1.5, 2.9, str(len(ho_venc)), 'já passaram da data', RED)
kpi(s, 9.7, 1.5, 3.05, '%dd' % HO[0]['antiga']['dias'],
    'a mais antiga (%s)' % usina_curta(HO[0]['u']), RED)


def fmt_ho(d):
    if d > 0:
        return fmt_dias(d)
    if -4 <= d <= 0:
        return ('vence esta semana', AMBER, False)  # C6 (verificado no JSON)
    return ('no prazo', GREEN, False)


linhas = [((usina_curta(h['u']), NAVY, True),
           ('OS %s' % ' / '.join(h['oss']), NAVY, False),
           str(h['tarefas']),
           h['antiga']['t'].replace('Handover – ', '').replace('Handover - ', '')
           .replace(' (Seco/Óleo)', '')[:40],
           h['antiga']['d'],
           fmt_ho(h['antiga']['dias']))
          for h in HO]
tabela(s, 0.55, 2.85, 12.2,
       ['Usina', 'OS', 'Tarefas', 'Escopo mais antigo', 'Programada', 'Idade'],
       [2.9, 1.35, 0.85, 4.1, 1.2, 1.8], linhas, hrow=0.222, sz=9,
       centro={2})  # V12
# C11: nota da OS 12876
tx(s, 4.7, 7.08, 8.05, 0.3,
   'OS 12876 atende duas classificações de Santo Antônio da Platina — '
   '17 linhas para 16 OS.', 8.5, BODY, False, PP_ALIGN.RIGHT)

# ── 10 BACKLOG +30D POR SUPERVISOR ───────────────────────────────────
# V3: título com 69 OS e 23 corretivas
s = slide('O estoque antigo por supervisor — %d OS há +30 dias (%d corretivas)'
          % (G['backlog_os'], G['corr_30d']), 'PCM - Backlog')
BL = D['backlog']
grupos = defaultdict(list)
for x in BL:
    grupos[x.get('sup') or 'Sem cluster'].append(x)
ordem = sorted(grupos.items(), key=lambda kv: -len(kv[1]))
col_y = {0.55: 1.45, 6.78: 1.45}
CHIP_W, CHIP_H, HEAD_H = 2.93, 0.21, 0.27
for sup, oss in ordem:
    oss.sort(key=lambda x: -x['dias'])
    alt = HEAD_H + 0.03 + ((len(oss) + 1) // 2) * CHIP_H + 0.06
    x0 = min(col_y, key=lambda k: col_y[k])
    y0 = col_y[x0]
    caixa(s, x0, y0, 5.95, HEAD_H - 0.02, NAVY, 0.16)
    tx(s, x0 + 0.14, y0 + 0.028, 4.4, 0.22, sup, 10.5, WHITE, True)
    tx(s, x0 + 4.3, y0 + 0.028, 1.5, 0.22, '%d OS' % len(oss), 10.5, LIME, True,
       PP_ALIGN.RIGHT)
    for i, x in enumerate(oss):
        cx = x0 + (i % 2) * (CHIP_W + 0.09)
        cy = y0 + HEAD_H + 0.03 + (i // 2) * CHIP_H
        cor = RED if x['dias'] > 60 else AMBER
        caixa(s, cx, cy, CHIP_W, 0.185, CALL, 0.2)
        # V5: barra lateral na cor da idade + nome da usina 8.5pt em INK
        retang(s, cx, cy + 0.025, 0.045, 0.135, cor)
        tx(s, cx + 0.12, cy + 0.018, CHIP_W - 0.18, 0.16,
           [[('OS %s ' % x['os'], NAVY, True, 8.5),
             ('%dd ' % x['dias'], cor, True, 8.5),
             ('· %s' % usina_curta(x['u'])[:19], INK, False, 8.5)]], 8.5)
    col_y[x0] = y0 + alt

# ── 11 DESAFIOS DA SEMANA ────────────────────────────────────────────
s = slide('Nossos desafios — o combinado da segunda', 'PCM - Semana 39')
desafios = [
    ('Fechar a MPM de setembro (100% até 30/09)',
     'o mês fecha na quarta — foco em RenoGrid, GD Energy e Alves Lima (0–20%%) e nas %d usinas pendentes da Thopen.'
     % len(tho['mpm_pend_usinas'])),
    # C2 + C3: donos e desligamento remarcado
    ('Zerar as MPA vencidas mais antigas',
     'Petrolina 2 (Axis, 104d, Vitor Valadares) — desligamento remarcado para 10/09: '
     'confirmar se ocorreu e se os ensaios foram concluídos; Mandaguaçu 1 (81d, Pedro '
     'Candido), Cedro 1/2 (79d, Marcelo Martineli) e as 5 da RenoGrid (Muito Crítico, '
     'Fred Alexandrino).'),
    ('Decidir as %d MPS vencidas da Thopen' % G['mps_venc'],
     'a 6340 fechou — agora a mais antiga é a SPDA de Araçoiaba da Serra 1 (81d, '
     'Danuth Fernandes): faz, reprograma ou cancela.'),
    # C6: não deixar vencer os handovers desta semana (verificado no JSON)
    ('Destravar os handovers mais antigos',
     'Coração 1 e 2 (65d), Colorado 1 (45d) e Diamantino 1 (41d) — plano por usina no '
     'slide 9 — e não deixar vencer os desta semana (Canarana 2, Alvares Machado, '
     'Assis Chateaubriand).'),
    # V3 + C3 + C4
    ('Reduzir o estoque de corretivas +30d',
     '%d corretivas dentro das %d OS +30d do slide 10 — começar pelas %d da Athon '
     '(Timon 1: trackers parados há %dd, Pedro Candido); cada supervisor olha o seu '
     'bloco e definir dono para as 7 OS "Sem cluster" (Marabá 2, Guajirú 1, Poconé 1) '
     'ainda hoje.'
     % (G['corr_30d'], G['backlog_os'], ath['corr_30d'], ath['antiga']['dias'])),
    ('Análise de óleo',
     '%d OS abertas — atualizar o estágio logístico (enviar / rota / usina / laboratório) na OS.' % G['oleo_ab']),
    ('Fechamento na sexta',
     'na sexta (25/09) revisamos estes pontos com os dados novos da API — o que fechou, o que ficou.'),
]
y = 1.5
for i, (t_, d_) in enumerate(desafios, 1):
    caixa(s, 0.55, y, 0.4, 0.4, LIME, 0.5)
    tx(s, 0.55, y + 0.075, 0.4, 0.28, str(i), 13, NAVY, True, PP_ALIGN.CENTER)  # V14
    tx(s, 1.12, y, 11.5, 0.72,
       [[(t_ + ' — ', INK, True, 12.5)], [(d_, BODY, False, 11)]], 12.5)
    y += 0.78

DEST = (r'C:\Users\Fabricio Barreto\OneDrive - GRID CO\Shortcuts'
        r'\GRID CO_ - 4. O&M\11.Pré-Operação\9. Reuniões\0. Grid Interno'
        r'\1. Inicio e Fechamento da Semana\Inicio Semana39_Desafios_PCM_R01.pptx')
prs.save('_s39_r01.pptx')
import time
for tent in range(6):
    try:
        prs.save(DEST)
        print('ok -> %s (%d slides)' % (DEST, len(prs.slides._sldIdLst)))
        break
    except PermissionError:
        if tent == 5:
            print('AVISO: destino OneDrive bloqueado (arquivo aberto?); '
                  'local _s39_r01.pptx salvo com %d slides.'
                  % len(prs.slides._sldIdLst))
        else:
            time.sleep(5)
