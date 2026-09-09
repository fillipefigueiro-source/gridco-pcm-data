# -*- coding: utf-8 -*-
"""
Exporta o universo de usinas do BD_Operacoes.xlsx (OneDrive) para operacoes.json.

Por que existe: a cascata de disponibilidade do módulo Gerencial precisa do
universo em operação, da potência e do nº de cabines e inversores por usina —
o peso de cada evento (cabine = 1/nº de cabines, inversor = 1/nº de inversores).
Isso vive no BD_Operacoes, que a nuvem não enxerga. Como o Gerencial MPAS:
roda na máquina do PCM quando o BD muda, e o JSON vai no commit.

Uso:  python exportar_operacoes.py [caminho\\BD_Operacoes.xlsx]
"""
from __future__ import annotations

import io
import json
import os
import pathlib
import re
import sys
from datetime import datetime

import openpyxl

CANDIDATOS = [
    os.environ.get("BD_OPERACOES_PATH", ""),
    str(pathlib.Path.home() / "GRID CO" / "Grid Co. - Gridco" / "4. O&M" / "6.Gerencial"
        / "4. Gestão à vista" / "1. Banco de Dados" / "BD_Operacoes.xlsx"),
]
SAIDA = os.environ.get("OPERACOES_SAIDA", "operacoes.json")

REGIAO_UF = {
    "AC": "Norte", "AM": "Norte", "AP": "Norte", "PA": "Norte", "RO": "Norte", "RR": "Norte", "TO": "Norte",
    "AL": "Nordeste", "BA": "Nordeste", "CE": "Nordeste", "MA": "Nordeste", "PB": "Nordeste", "PE": "Nordeste",
    "PI": "Nordeste", "RN": "Nordeste", "SE": "Nordeste",
    "DF": "Centro-Oeste", "GO": "Centro-Oeste", "MT": "Centro-Oeste", "MS": "Centro-Oeste",
    "ES": "Sudeste", "MG": "Sudeste", "RJ": "Sudeste", "SP": "Sudeste",
    "PR": "Sul", "RS": "Sul", "SC": "Sul",
}


def _num(v):
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "."))
    except Exception:
        return None


def _int(v):
    n = _num(v)
    return int(round(n)) if n is not None else None


def _col(cab, *nomes):
    """Índice da coluna cujo cabeçalho começa com um dos nomes (sem caixa)."""
    for i, c in enumerate(cab):
        s = str(c or "").strip().lower()
        for n in nomes:
            if s.startswith(n.lower()):
                return i
    return None


def main(argv):
    caminho = next((p for p in ([argv[1]] if len(argv) > 1 else []) + CANDIDATOS if p and os.path.exists(p)), None)
    if not caminho:
        print("BD_Operacoes.xlsx não encontrado — passe o caminho ou defina BD_OPERACOES_PATH")
        return 2
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb["Operações"]
    linhas = ws.iter_rows(values_only=True)
    cab = next(linhas)
    ix = {
        "usina": _col(cab, "OPERAÇÃO"), "codigo": _col(cab, "CÓDIGO"), "status": _col(cab, "STATUS"),
        "cliente": _col(cab, "CLIENTE"), "mwp": _col(cab, "POTÊNCIA REAL"), "mwpContr": _col(cab, "POTÊNCIA CONTRATUAL"),
        "uf": _col(cab, "UF"), "nInv": _col(cab, "Nº INVERSORES"), "nCab": _col(cab, "Nº Cabine"),
        "nSkid": _col(cab, "Nº SKID"), "nQgbt": _col(cab, "Nº QGBT"), "cidade": _col(cab, "CIDADE"),
        "implantacao": _col(cab, "Data de Implantação"),
        # cluster, região oficial e responsável — o BD tem; a UF é só o recuo
        # "Equipe Cluster" (PR Oeste 01) é o nome que a aba de colaboradores usa;
        # "CLUSTER" (PR Oeste) é a região do cluster, sem o número — não casa.
        "cluster": _col(cab, "Equipe Cluster"), "clusterAlt": _col(cab, "CLUSTER"),
        "regiaoBd": _col(cab, "REGIÃO"), "responsavel": _col(cab, "RESPONSÁVEL O&M"),
    }
    faltando = [k for k, v in ix.items() if v is None and k in ("usina", "codigo", "status", "cliente", "uf")]
    if faltando:
        print("colunas obrigatórias não encontradas:", faltando)
        return 2
    out = []
    for r in linhas:
        if not r or not r[ix["usina"]]:
            continue
        g = lambda k: (r[ix[k]] if ix[k] is not None else None)
        uf = str(g("uf") or "").strip().upper()[:2]
        impl = g("implantacao")
        out.append({
            "usina": str(g("usina")).strip(),
            "codigo": str(g("codigo") or "").strip(),
            "status": str(g("status") or "").strip().upper(),
            "cliente": str(g("cliente") or "").strip(),
            "uf": uf, "regiao": (re.sub(r"centro[\s-]*oeste", "Centro-Oeste", str(g("regiaoBd") or "").strip().title(), flags=re.I) or REGIAO_UF.get(uf, "")),
            "cluster": str(g("cluster") or g("clusterAlt") or "").strip(), "responsavel": str(g("responsavel") or "").strip(),
            "cidade": str(g("cidade") or "").strip(),
            "mwp": _num(g("mwp")) or _num(g("mwpContr")),
            "nInv": _int(g("nInv")), "nCab": _int(g("nCab")), "nSkid": _int(g("nSkid")), "nQgbt": _int(g("nQgbt")),
            "implantacao": impl.strftime("%Y-%m-%d") if isinstance(impl, datetime) else (str(impl)[:10] if impl else ""),
        })
    oper = [u for u in out if u["status"] == "OPERAÇÃO"]
    semReg = [u["usina"] for u in oper if not u["regiao"]]
    semInv = [u["usina"] for u in oper if not u["nInv"]]

    # ── Equipe: aba "Relação Geral Colaboradores" ──────────────────────────
    # Só o que a plataforma precisa: nome, cluster, cargo, e-mail, supervisor.
    # CPF, telefones e endereço NÃO saem daqui — o repositório é público.
    equipe = []
    if "Relação Geral Colaboradores" in wb.sheetnames:
        wc = wb["Relação Geral Colaboradores"]
        lc = wc.iter_rows(values_only=True)
        cabc = next(lc)
        ic = {k: _col(cabc, n) for k, n in (("nome", "Nome"), ("padrao", "Nome Padrão"), ("cluster", "Cluster"), ("cargo", "Cargo"),
                                            ("email", "email"), ("sup", "Supervisor"), ("cliente", "Cliente"), ("status", "Status de Contratação"))}
        for r in lc:
            if not r or ic["nome"] is None or not r[ic["nome"]]:
                continue
            gc = lambda k: (r[ic[k]] if ic[k] is not None else None)
            if str(gc("status") or "").strip().lower() != "ativo":
                continue
            email = str(gc("email") or "").strip().lower()
            equipe.append({
                "nome": str(gc("padrao") or gc("nome")).strip(), "nomeCompleto": str(gc("nome")).strip(),
                "cluster": str(gc("cluster") or "").strip(), "cargo": str(gc("cargo") or "").strip(),
                "email": email if "@" in email else "", "supervisor": str(gc("sup") or "").strip(),
                "cliente": str(gc("cliente") or "").strip(),
            })
    # HH: cada colaborador tem 44 h/semana; o cluster soma as pessoas, e o HH
    # por usina é o do cluster dividido pelas usinas do cluster (regra do PCM, 09/09).
    HH_SEMANA = 44
    clusters = {}
    for u in oper:
        c = u["cluster"] or "(sem cluster)"
        clusters.setdefault(c, {"cluster": c, "usinas": 0, "pessoas": 0, "hhSemana": 0, "hhPorUsina": None, "responsavel": u["responsavel"]})
        clusters[c]["usinas"] += 1
    for e in equipe:
        c = e["cluster"] or "(sem cluster)"
        clusters.setdefault(c, {"cluster": c, "usinas": 0, "pessoas": 0, "hhSemana": 0, "hhPorUsina": None, "responsavel": ""})
        clusters[c]["pessoas"] += 1
        clusters[c]["hhSemana"] += HH_SEMANA
    for c in clusters.values():
        c["hhPorUsina"] = round(c["hhSemana"] / c["usinas"], 1) if c["usinas"] else None
    supervisores = sorted({e["supervisor"] for e in equipe if e["supervisor"]})

    saida = {
        "geradoEm": datetime.now().replace(microsecond=0).isoformat(),
        "fonte": os.path.basename(caminho), "fonteModificadoEm": datetime.fromtimestamp(os.path.getmtime(caminho)).isoformat(timespec="seconds"),
        "total": len(out), "emOperacao": len(oper),
        "mwpOperacao": round(sum(u["mwp"] or 0 for u in oper), 1),
        "usinas": out,
        "equipe": sorted(equipe, key=lambda e: e["nome"]), "supervisores": supervisores,
        "hhSemanaPorPessoa": HH_SEMANA, "clusters": sorted(clusters.values(), key=lambda c: c["cluster"]),
    }
    with io.open(SAIDA, "w", encoding="utf-8") as f:
        json.dump(saida, f, ensure_ascii=False, separators=(",", ":"))
    print(f"{SAIDA}: {len(out)} usinas, {len(oper)} em OPERAÇÃO, {saida['mwpOperacao']} MWp"
          + (f" | SEM REGIÃO: {semReg}" if semReg else "") + (f" | sem nº de inversores: {len(semInv)}" if semInv else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
