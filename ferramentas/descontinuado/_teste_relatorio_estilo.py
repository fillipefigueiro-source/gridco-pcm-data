# -*- coding: utf-8 -*-
"""
TESTE / EXPERIMENTAL — relatório de PCM no ESTILO do relatório de Performance UFV.
NÃO é o sistema oficial. Reusa os CÁLCULOS de relatorio_clientes.py (import) e só
renderiza num layout diferente (capa com veredito + seções numeradas com título-pergunta
+ tiles de KPI + veredito em caixa alta + pontos de atenção + tabelas + gráficos).

Uso:
    python _teste_relatorio_estilo.py --usina "Axis - Petrolina 2" --ref 2026-06
Saída: Relatórios/_Teste_Estilo/TESTE_Estilo_<usina>_<mes>.pdf
"""
import os, sys, argparse, datetime as dt
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

import relatorio_clientes as R
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# Atalhos p/ a identidade visual e helpers já existentes no engine oficial
NAVY, GRAY, LIME, WHITE, LIGHT, MID, INK = R.NAVY, R.GRAY, R.LIME, R.WHITE, R.LIGHT, R.MID, R.INK
ST_G, ST_Y, ST_R = R.ST_G, R.ST_Y, R.ST_R
hx = R.hx

# ---------------------------------------------------------------- CAPA (estilo período)
def capa_periodo(nome, cliente, cap, rot, verdict_txt, verdict_cor, headline, resumo, fn):
    fig = plt.figure(figsize=(8.27, 11.69), dpi=170); ax = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
    ax.add_patch(plt.Rectangle((0, 0), 1, 1, color="#" + NAVY))
    ax.add_patch(plt.Rectangle((0.0, 0.0), 0.018, 1, color="#" + LIME))
    try:
        import matplotlib.image as mpimg
        em = mpimg.imread(os.path.join(R.ASSETS, "emblema_white.png"))
        axw = fig.add_axes([0.40, -0.05, 0.72, 0.50]); axw.imshow(em, alpha=0.05); axw.axis("off")
        em2 = mpimg.imread(os.path.join(R.ASSETS, "emblema_lime.png"))
        axe = fig.add_axes([0.10, 0.86, 0.11, 0.11]); axe.imshow(em2); axe.axis("off")
    except Exception:
        pass
    T = ax.transAxes
    ax.text(0.225, 0.90, "Grid Co.", fontproperties=R.fp(R.FB, 22), color="#FFFFFF", transform=T)
    ax.text(0.225, 0.875, "OPERAÇÃO & MANUTENÇÃO", fontproperties=R.fp(R.FR, 9.5), color="#" + LIME, transform=T)
    # faixa título
    ax.text(0.10, 0.70, "RELATÓRIO DE MANUTENÇÃO", fontproperties=R.fp(R.FR, 12), color="#" + MID, transform=T)
    ax.text(0.10, 0.655, "RELATÓRIO DO PERÍODO  ·  " + rot.upper(), fontproperties=R.fp(R.FR, 9), color="#8A8794", transform=T)
    ax.text(0.10, 0.585, nome, fontproperties=R.fp(R.FB, 30), color="#FFFFFF", transform=T)
    sub = f"{R.fnum(cap,2)} MWp  ·  {cliente}" if cap else str(cliente)
    ax.text(0.10, 0.545, sub, fontproperties=R.fp(R.FS, 13), color="#" + LIME, transform=T)
    # selo de veredito (pill)
    ax.add_patch(plt.Rectangle((0.10, 0.455), 0.44, 0.055, color="#" + verdict_cor, transform=T, zorder=2))
    ax.text(0.12, 0.472, "●  " + verdict_txt, fontproperties=R.fp(R.FS, 13), color="#FFFFFF", transform=T, zorder=3)
    # headline (2 números de destaque)
    y = 0.36
    for val, lab in headline:
        ax.text(0.10, y, val, fontproperties=R.fp(R.FB, 26), color="#FFFFFF", transform=T)
        ax.text(0.335, y + 0.012, lab, fontproperties=R.fp(R.FR, 11), color="#" + MID, transform=T)
        y -= 0.075
    # resumo
    import textwrap
    yy = 0.175
    for ln in textwrap.wrap(resumo, 82):
        ax.text(0.10, yy, ln, fontproperties=R.fp(R.FR, 10.5), color="#D8D6DE", transform=T); yy -= 0.028
    ax.text(0.10, 0.055, "Equipe de PCM  ·  Grid Co.", fontproperties=R.fp(R.FS, 10), color="#" + LIME, transform=T)
    ax.text(0.10, 0.03, dt.date.today().strftime("Emitido em %d/%m/%Y"), fontproperties=R.fp(R.FR, 8.5), color="#8A8794", transform=T)
    p = os.path.join(R.CHARTS, fn); fig.savefig(p, dpi=170); plt.close(fig); return p

# ---------------------------------------------------------------- helpers de layout (estilo período)
def secao(doc, num, titulo, pergunta, explica=None, page_break=True):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    r0 = p.add_run(f"{num}  "); r0.font.name = "Poppins"; r0.font.bold = True; r0.font.size = Pt(15); r0.font.color.rgb = hx(LIME)
    r1 = p.add_run(titulo.upper()); r1.font.name = "Poppins"; r1.font.bold = True; r1.font.size = Pt(15); r1.font.color.rgb = hx(NAVY)
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(1)
    pq = doc.add_paragraph(); rq = pq.add_run(pergunta); rq.font.name = "Poppins"; rq.font.italic = True; rq.font.size = Pt(12.5); rq.font.color.rgb = hx(GRAY)
    pq.paragraph_format.space_after = Pt(4)
    if explica:
        R.body(doc, explica, 9.5, MID)

def veredito(doc, texto, cor=NAVY):
    p = doc.add_paragraph(); r = p.add_run(texto.upper()); r.font.name = "Poppins"; r.font.bold = True; r.font.size = Pt(11); r.font.color.rgb = hx(cor)
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
    R._pbottom(p, LIME, 8)

def kpi_tiles(doc, items, per_row=3):
    """items = list of (valor, rotulo, referencia)."""
    for i in range(0, len(items), per_row):
        bloco = items[i:i + per_row]
        t = doc.add_table(rows=1, cols=len(bloco)); R.no_borders(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, (val, rot, ref) in enumerate(bloco):
            c = t.cell(0, j); R.shade(c, LIGHT); R.cmarg(c, 120, 120, 140, 140); c.vertical_alignment = 1
            p = c.paragraphs[0]; r = p.add_run(val); r.font.name = "Poppins"; r.font.bold = True; r.font.size = Pt(20); r.font.color.rgb = hx(NAVY)
            p2 = c.add_paragraph(); r2 = p2.add_run(rot); r2.font.name = "Poppins"; r2.font.size = Pt(8.5); r2.font.color.rgb = hx(GRAY)
            p3 = c.add_paragraph(); r3 = p3.add_run(ref); r3.font.name = "Poppins"; r3.font.size = Pt(7.5); r3.font.color.rgb = hx(MID)
            p2.paragraph_format.space_before = Pt(1); p3.paragraph_format.space_before = Pt(0)
            R._cell_bottom(c, LIME, 14)
        doc.add_paragraph()

def ponto(doc, marker, texto, cor):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.1)
    b = p.add_run(marker + "  "); b.font.name = "Poppins"; b.font.bold = True; b.font.size = Pt(10); b.font.color.rgb = hx(cor)
    r = p.add_run(texto); r.font.size = Pt(10); r.font.color.rgb = hx(INK); p.paragraph_format.space_after = Pt(2)

def tabela(doc, cols, linhas):
    t = doc.add_table(rows=1, cols=len(cols)); t.style = "Table Grid"; R.th_row(t, cols)
    for row in linhas:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            R.td(cells[j], v, 8.5, bold=(str(row[0]).upper() == "TOTAL"))
    return t

def fmt(v, d=0, suf=""):
    try:
        return R.fnum(v, d) + suf
    except Exception:
        return str(v) + suf

# ---------------------------------------------------------------- MONTAGEM
def gerar(usina, ref):
    idx, rows, aux, cu = R.carregar(".")
    regs = R.coletar("usina", usina, idx, rows, aux)
    if not regs:
        print(f"[ERRO] nenhum registro para a usina '{usina}'. Confira o nome exato (Ativo Classificação 1).")
        return None
    ini, fim, rot, tipo = R.janela("mensal", ref)
    k = R.kpis_mes(regs, ini, fim); conc = k["conc"]
    mpm = R.mp_mes(regs, ini, fim); mp12c, _ = R.mp_12_por_code(regs, fim)
    hist = R.historico_mensal(regs, fim)
    tipos = R.confiab_tipologia(regs, fim); crit_list = R.crit_corretivas(conc)
    ativos_conf = R.confiab_ativo(regs, ini, fim); ader = R.aderencia_prev(regs, ini, fim)
    tempos = R.tempos_atend(conc)
    n_us = 1
    disp_media = (sum(t["disp"] * t["corr"] for t in tipos) / sum(t["corr"] for t in tipos)) if tipos else 100.0
    mttr_medio = (sum(t["mttr"] * t["corr"] for t in tipos) / sum(t["corr"] for t in tipos)) if any(t["corr"] for t in tipos) else 0.0
    mtbf_medio = (sum(t["mtbf"] * t["corr"] for t in tipos) / sum(t["corr"] for t in tipos)) if any(t["corr"] for t in tipos) else 0.0
    cliente = max([r["cli"] for r in regs if r.get("cli")], default="—", key=[r["cli"] for r in regs].count) if regs else "—"
    cap = next((r.get("cap") for r in regs if r.get("cap")), 0) or 0
    emerg = k["emerg"]
    crit_alta = sum(v for n, v in crit_list if n in ("Muito alto", "Alto"))
    top_ativo = k["corr_ativo"].most_common(1)[0] if k["corr_ativo"] else None

    # ---- veredito
    avanco = mpm["TOTAL"]["pct"]; pctprev = k["pct_prev_os"]
    if avanco >= 75 and pctprev >= 70:
        v_txt, v_cor = "Operação sob controle", ST_G
    elif avanco >= 50 or pctprev >= 50:
        v_txt, v_cor = "Requer atenção", ST_Y
    else:
        v_txt, v_cor = "Ação corretiva necessária", ST_R

    resumo = (f"No período, {usina} registrou {k['n_os']} ordens de trabalho "
              f"({k['nprev_os']} preventivas, {k['ncorr_os']} corretivas) — {pctprev:.0f}% de preventiva. "
              f"Avanço do plano preventivo do mês em {avanco:.0f}%. {k['abert_os']} tarefa(s) em aberto.")
    headline = [(f"{pctprev:.0f}%", "de manutenção preventiva"),
                (f"{avanco:.0f}%", "de avanço do plano do mês")]
    capa = capa_periodo(usina, cliente, cap, rot, v_txt, v_cor, headline, resumo, "capa_teste.png")

    # ---- gráficos reaproveitados do engine
    gh = R.g_hist(hist, "hist_teste.png"); gano = R.g_corr_ano(hist, "corr_teste.png")

    # ==================== DOCUMENTO ====================
    doc = Document(); stl = doc.styles["Normal"]; stl.font.name = "Calibri"; stl.font.size = Pt(10.5)
    sec = doc.sections[0]
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(0))
    doc.add_picture(capa, width=Inches(8.27))
    doc.add_section(); s2 = doc.sections[1]
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(s2, m, Inches(0.85))
    s2.top_margin = Inches(0.95); R.footer(s2)
    R.header_setup(s2, "Usina " + usina, rot, width_in=6.5)

    # 01 RESUMO EXECUTIVO
    secao(doc, "01", "Resumo Executivo", "O período em uma página",
          "Indicadores essenciais de manutenção medidos contra as referências de O&M, o veredito e os pontos que exigem atenção.",
          page_break=False)
    kpi_tiles(doc, [
        (f"{pctprev:.0f}%", "% Preventiva", "referência ≥ 70%"),
        (f"{avanco:.0f}%", "Avanço do Plano (mês)", "referência ≥ 90%"),
        (f"{k['pct_prog_os']:.0f}%", "% Programadas", "referência ≥ 60%"),
        (f"{disp_media:.1f}%", "Disponib. Inerente", "referência ≥ 98%"),
        (fmt(mttr_medio, 1, " h"), "MTTR corretiva", "menor = melhor"),
        (str(k["abert_os"]), "Tarefas em Aberto", f"de {k['n_os']} no período"),
    ])
    veredito(doc, f"Veredito: {v_txt}", v_cor)
    R.body(doc, resumo)
    doc.add_paragraph()
    pat = doc.add_paragraph(); rr = pat.add_run("Pontos de atenção e ações prioritárias"); rr.font.name = "Poppins"; rr.font.bold = True; rr.font.size = Pt(10.5); rr.font.color.rgb = hx(NAVY)
    if crit_alta:
        ponto(doc, "!", f"{crit_alta} corretiva(s) de criticidade alta/muito alta no período — priorizar na escala e no estoque de sobressalentes.", ST_Y)
    if top_ativo:
        ponto(doc, "!", f"Maior incidência corretiva: '{str(top_ativo[0])[:52]}' ({top_ativo[1]} corretiva(s)).", ST_Y)
    if k["abert_os"]:
        ponto(doc, "!", f"{k['abert_os']} tarefa(s) ainda em aberto — acompanhar vazão do backlog.", ST_Y)
    if not (crit_alta or (top_ativo and top_ativo[1] > 1) or k["abert_os"]):
        ponto(doc, "✓", "Sem ofensores relevantes no período. Operação dentro do esperado.", ST_G)

    # 02 DEMANDA DE MANUTENÇÃO
    secao(doc, "02", "Demanda de Manutenção", "Quanto trabalho apareceu — e quanto viramos?",
          "Volume de ordens no período e a relação preventiva × corretiva; ao lado, a evolução no ano.")
    kpi_tiles(doc, [
        (str(k["n_os"]), "OTs no período", "total"),
        (str(k["n_os"] - k["abert_os"]), "Concluídas", f"{100*(k['n_os']-k['abert_os'])//max(1,k['n_os'])}%"),
        (str(k["abert_os"]), "Em aberto", "backlog"),
        (str(k["nprev_os"]), "Preventivas", f"{pctprev:.0f}% do total"),
        (str(k["ncorr_os"]), "Corretivas", f"{100-pctprev:.0f}% do total"),
        (str(emerg), "Emergenciais", "corretiva urgente"),
    ])
    R.add_img(doc, gh, 6.4)
    veredito(doc, "Leitura", NAVY)
    R.body(doc, f"A relação preventiva × corretiva alcançou {pctprev:.0f}% de preventiva, "
                f"{'acima' if pctprev>=70 else 'abaixo'} da referência de 70% para um regime confiável. "
                f"{k['prog_os']} tarefa(s) foram programadas ({k['pct_prog_os']:.0f}%), refletindo o grau de planejamento da operação.")

    # 03 CONFIABILIDADE DOS ATIVOS
    secao(doc, "03", "Confiabilidade dos Ativos", "Os ativos estão confiáveis?",
          "MTTR, MTBF e disponibilidade inerente por tipologia de ativo, desde a mobilização.")
    kpi_tiles(doc, [
        (fmt(mttr_medio, 1, " h"), "MTTR médio", "tempo de reparo"),
        (fmt(mtbf_medio, 0, " h"), "MTBF médio", "entre falhas"),
        (f"{disp_media:.1f}%", "Disponib. Inerente", "MTBF ÷ (MTBF+MTTR)"),
    ])
    if tipos:
        tabela(doc, ["Tipologia de Ativo", "Corretivas", "MTTR (h)", "MTBF (h)", "Disponib. (%)"],
               [[t["tip"], t["corr"], f"{t['mttr']:.1f}", fmt(t["mtbf"], 0), f"{t['disp']:.1f}%"] for t in tipos])
        pf = tipos[0]; pm = max(tipos, key=lambda t: t["mttr"])
        veredito(doc, "Leitura", NAVY)
        R.body(doc, f"'{pf['tip']}' lidera em corretivas ({pf['corr']}); '{pm['tip']}' tem o maior MTTR "
                    f"({pm['mttr']:.1f} h) — foco em sobressalentes e capacitação para reduzir o tempo de reparo.")
    else:
        R.body(doc, "Sem corretivas com tempos apurados no período para cálculo de confiabilidade.")

    # 04 PLANO PREVENTIVO
    secao(doc, "04", "Plano Preventivo", "O plano saiu do papel?",
          "Preventivas programadas para o período e quantas foram finalizadas, por periodicidade.")
    kpi_tiles(doc, [
        (f"{mpm['TOTAL']['pct']:.0f}%", "Avanço do plano (mês)", f"{mpm['TOTAL']['fin']}/{mpm['TOTAL']['plan']}"),
        (f"{ader['pct_prazo']:.0f}%", "Aderência ao prazo", f"{ader['no_prazo']}/{ader['fin']} no prazo"),
        (str(mpm["TOTAL"]["pend"]), "Pendentes no mês", "a concluir"),
    ])
    codes = [c for c in R.MP_ORDER if mpm[c]["plan"] > 0]
    if codes:
        tabela(doc, ["Periodicidade", "Programadas", "Finalizadas", "Pendentes", "% Avanço"],
               [[R.MP_LABEL.get(c, c), mpm[c]["plan"], mpm[c]["fin"], mpm[c]["pend"], f"{mpm[c]['pct']:.0f}%"] for c in codes] +
               [["TOTAL", mpm["TOTAL"]["plan"], mpm["TOTAL"]["fin"], mpm["TOTAL"]["pend"], f"{mpm['TOTAL']['pct']:.0f}%"]])
    veredito(doc, "Leitura", NAVY)
    R.body(doc, f"Avanço geral de {mpm['TOTAL']['pct']:.0f}% no mês ({mpm['TOTAL']['fin']}/{mpm['TOTAL']['plan']} tarefas). "
                f"Aderência ao cronograma: {ader['pct_prazo']:.0f}% das finalizadas concluídas até a data programada.")

    # 05 CORRETIVAS & OFENSORES
    secao(doc, "05", "Corretivas & Ofensores", "O que quebrou — e por quê?",
          "Corretivas do período, criticidade e os ativos que mais demandaram intervenção.")
    kpi_tiles(doc, [
        (str(k["ncorr_os"]), "Corretivas", "no período"),
        (str(emerg), "Emergenciais", "urgentes"),
        (str(crit_alta), "Criticidade alta", "alta/muito alta"),
    ])
    top = sorted([d for d in ativos_conf if d["corr"] > 0], key=lambda d: d["corr"], reverse=True)[:10]
    if top:
        tabela(doc, ["Ativo / Equipamento", "Tipologia", "Corretivas", "Δ vs mês ant."],
               [[str(d["ativo"])[:42], R.tipologia_ativo(d["ativo"]), d["corr"],
                 ("▲ +" + str(d["delta"])) if d["delta"] > 0 else (("▼ " + str(d["delta"])) if d["delta"] < 0 else "—")] for d in top])
    veredito(doc, "Leitura", NAVY)
    R.body(doc, (f"Foram {k['ncorr_os']} corretiva(s) no período" +
                 (f", sendo {crit_alta} de criticidade alta/muito alta. " if crit_alta else ". ") +
                 (f"O ativo com maior incidência foi '{str(top_ativo[0])[:48]}'." if top_ativo else "Sem concentração relevante em um único ativo.")))

    # 06 EVOLUÇÃO NO ANO
    secao(doc, "06", "Evolução no Ano", "A visão acumulada do ano",
          "A série mês a mês da demanda de manutenção e da relação preventiva × corretiva.")
    R.add_img(doc, gh, 6.4)
    R.add_img(doc, gano, 6.4)
    veredito(doc, "Leitura", NAVY)
    prev_ini = hist[0]["pct"] if hist else 0; prev_fim = hist[-1]["pct"] if hist else 0
    R.body(doc, f"A parcela de preventiva saiu de {prev_ini:.0f}% em {hist[0]['rot']} para {prev_fim:.0f}% em {hist[-1]['rot']}. "
                "Uma tendência de alta na preventiva indica amadurecimento do plano e menor reatividade.")

    # 07 PENDÊNCIAS & PRÓXIMOS PASSOS
    secao(doc, "07", "Pendências & Próximos Passos", "O que precisa de ação?",
          "Preventivas pendentes no mês e recomendações para o próximo ciclo.")
    pend = [(c, x) for c in R.MP_ORDER for x in mpm[c]["pendentes"][:5]]
    if pend:
        for c, x in pend[:10]:
            ponto(doc, "•", f"{c} · {str(x['ufv']).split(' - ')[-1][:22]} · {str(x['tarefa'])[:44]} (prog. {R.fdate(x['dprog'])})", GRAY)
    else:
        ponto(doc, "✓", "Sem preventivas pendentes no período. Plano do mês concluído.", ST_G)
    if k["abert_os"]:
        ponto(doc, "!", f"{k['abert_os']} tarefa(s) em aberto para vazão no próximo ciclo.", ST_Y)

    # 08 METODOLOGIA
    secao(doc, "08", "Metodologia & Fórmulas", "Como os indicadores são calculados",
          "Definição dos principais indicadores de manutenção usados neste relatório.")
    for tit, txt in [
        ("% Preventiva", "preventivas ÷ (preventivas + corretivas) × 100. Quanto maior, mais madura a operação."),
        ("Avanço do plano (mês)", "preventivas finalizadas ÷ preventivas programadas no mês × 100 (cronograma inicia na mobilização da UFV)."),
        ("% Programadas", "tarefas programadas (preventiva/preditiva/handover) ÷ total concluído × 100."),
        ("MTTR", "tempo médio de reparo das corretivas (evento → finalização). MTBF = tempo de operação ÷ nº de corretivas."),
        ("Disponibilidade inerente", "MTBF ÷ (MTBF + MTTR) × 100, por tipologia de ativo, desde a mobilização."),
        ("Aderência ao prazo", "preventivas finalizadas até a data programada ÷ total finalizado × 100."),
    ]:
        p = doc.add_paragraph(); a = p.add_run(tit + ": "); a.font.bold = True; a.font.size = Pt(9.5); a.font.color.rgb = hx(NAVY); a.font.name = "Poppins"
        b = p.add_run(txt); b.font.size = Pt(9.5); b.font.color.rgb = hx(INK); p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.2

    R.aplicar_estilo_tabelas(doc)

    # salvar
    out_dir = os.path.join("Relatórios", "_Teste_Estilo"); os.makedirs(out_dir, exist_ok=True)
    base = f"TESTE_Estilo_{usina.replace(' ', '_').replace('/', '-')}_{rot.replace(' ', '_')}"
    docx_path = os.path.join(out_dir, base + ".docx"); doc.save(docx_path)
    print(f"[OK] DOCX: {docx_path}")
    try:
        pdf = R.to_pdf(docx_path); print(f"[OK] PDF: {pdf}")
    except Exception as e:
        print(f"[aviso] conversão p/ PDF falhou ({e}); abra o DOCX.")
    return docx_path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--usina", default="Axis - Petrolina 2")
    ap.add_argument("--ref", default="2026-06")
    a = ap.parse_args()
    gerar(a.usina, a.ref)
