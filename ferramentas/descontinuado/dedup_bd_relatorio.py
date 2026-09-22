# -*- coding: utf-8 -*-
"""
dedup_bd_relatorio.py
---------------------
Deduplica linhas duplicadas no BD_Relatório Semanal.xlsx (resultado do bug
intermitente da paginação Fracttal — task #167).

Estratégia:
1. Lê o BD
2. Agrupa por (OSs ID, Código, Tarefa) — combo que define unicamente uma linha
3. Mantém só 1 cópia de cada
4. Salva (com backup)

Uso:
   py -3 dedup_bd_relatorio.py
   py -3 dedup_bd_relatorio.py --dry-run        # só mostra, não salva
"""
from __future__ import annotations
import os, sys, argparse, datetime as dt, shutil
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

import openpyxl


def _log(msg):
    print(f"[{dt.datetime.now().strftime('%H:%M:%S')}] {msg}", flush=True)


def _detectar_pasta() -> Path:
    env = os.getenv("PCM_PROG_DIR")
    if env and Path(env).exists():
        return Path(env)
    here = Path(__file__).resolve().parent
    if (here / "BD_Relatório Semanal.xlsx").exists():
        return here
    for c in [
        Path.home() / "GRID CO" / "Grid Co. - Gridco" / "4. O&M"
            / "11.Pré-Operação" / "6. PCM" / "09. Programação Semanal",
    ]:
        if c.exists():
            return c
    raise FileNotFoundError("Pasta 09. Programação Semanal não encontrada.")


def dedup_bd(path_bd: Path, dry_run: bool = False) -> int:
    """Deduplica BD. Retorna número de linhas removidas (positivo) ou -1 erro."""
    _log(f"Lendo {path_bd.name}...")
    wb = openpyxl.load_workbook(path_bd)
    abas = wb.sheetnames
    _log(f"   Abas: {abas}")

    # Aba principal
    ws_principal = wb["Semanal"] if "Semanal" in abas else wb.active
    _log(f"   Aba principal: '{ws_principal.title}' | {ws_principal.max_row} linhas x {ws_principal.max_column} cols")

    # Pega cabeçalho
    header_row = [c.value for c in ws_principal[1]]
    header = [str(c).strip() if c else "" for c in header_row]
    if "OSs ID" not in header:
        _log(f"ERRO: coluna 'OSs ID' não encontrada. Headers: {header[:15]}")
        return -1
    idx_os = header.index("OSs ID")
    idx_cod = header.index("Código") if "Código" in header else None
    idx_tar = header.index("Tarefa") if "Tarefa" in header else None

    _log(f"   Indices: OSs ID={idx_os}, Código={idx_cod}, Tarefa={idx_tar}")

    # Coleta todas as linhas (depois da header)
    todas_linhas = list(ws_principal.iter_rows(min_row=2, values_only=True))
    _log(f"   Total linhas de dados: {len(todas_linhas)}")

    # Dedup por (OSs ID, Código, Tarefa)
    vistos = set()
    linhas_unicas = []
    for row in todas_linhas:
        os_id = row[idx_os] if idx_os < len(row) else None
        cod = row[idx_cod] if idx_cod is not None and idx_cod < len(row) else None
        tar = row[idx_tar] if idx_tar is not None and idx_tar < len(row) else None
        key = (os_id, cod, tar)
        if key in vistos:
            continue
        vistos.add(key)
        linhas_unicas.append(row)

    removidas = len(todas_linhas) - len(linhas_unicas)
    _log(f"   Linhas únicas: {len(linhas_unicas)} | Removidas: {removidas}")

    # Conta OSs únicas
    oss_unicas = len(set(r[idx_os] for r in linhas_unicas if r[idx_os] not in (None, "", 0)))
    _log(f"   OSs únicas no resultado: {oss_unicas}")

    if dry_run:
        _log("DRY-RUN: nada foi salvo. Use sem --dry-run pra aplicar.")
        return removidas

    if removidas == 0:
        _log("Nenhuma duplicata encontrada. Nada a fazer.")
        return 0

    # Backup antes de mexer
    bak_dir = path_bd.parent / "_backup_BD"
    bak_dir.mkdir(exist_ok=True)
    ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    bak_path = bak_dir / f"{path_bd.stem}_PRE_DEDUP_{ts}{path_bd.suffix}"
    _log(f"Backup: {bak_path.name}")
    shutil.copy2(path_bd, bak_path)

    # Limpa aba e regrava
    _log("Regravando aba 'Semanal' deduplicada...")
    # Apaga linhas existentes (mantém header)
    ws_principal.delete_rows(2, ws_principal.max_row)
    # Insere linhas deduplicadas
    for row in linhas_unicas:
        ws_principal.append(row)

    _log("Salvando...")
    wb.save(path_bd)
    _log(f"OK: {removidas} linhas duplicadas removidas. BD agora tem {len(linhas_unicas)} linhas.")
    return removidas


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dry-run", action="store_true",
                        help="Apenas mostra o que mudaria, sem salvar.")
    parser.add_argument("--arquivo", default="BD_Relatório Semanal.xlsx")
    args = parser.parse_args()

    pasta = _detectar_pasta()
    _log(f"Pasta: {pasta}")
    path_bd = pasta / args.arquivo
    if not path_bd.exists():
        _log(f"ERRO: {path_bd} não existe.")
        return 2

    try:
        n = dedup_bd(path_bd, dry_run=args.dry_run)
        return 0 if n >= 0 else 4
    except Exception as e:
        import traceback
        _log(f"FALHA: {type(e).__name__}: {e}")
        _log(traceback.format_exc())
        return 3


if __name__ == "__main__":
    sys.exit(main())
