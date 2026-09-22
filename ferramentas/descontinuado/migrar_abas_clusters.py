"""
Migra arquivos Programação Semana XX.xlsx existentes para incluir uma aba
para cada cluster da Grid Co., mesmo que sem tarefas. Idempotente —
abas que já existem não são tocadas.

Uso:
    python migrar_abas_clusters.py          # migra todos os arquivos da pasta base
    python migrar_abas_clusters.py 23       # migra só a semana 23
    python migrar_abas_clusters.py 21 22 23 # migra várias semanas

A pasta base é detectada via PCM_PROG_DIR ou auto-detecção do OneDrive.
"""

import os
import re
import sys
import pathlib
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment


# Mesmas colunas do programacao_v7.py (output padrão de aba de equipe)
COLS_PADRAO_EQUIPE = [
    'Equipe', 'Dia', 'OSs ID', 'Ativo (Usina)', 'Responsável', 'Cidade',
    'Código Equipamento', 'Tipo', 'Tarefa', 'Estado Tarefa (antes)',
    'RPN/Prioridade', 'Reprogramada', 'Nº vezes programada', 'Termografia',
    'Hora Início', 'Hora Fim', 'Duração (h)', 'Desloc (h)',
    'Paralelo (terceirizada)',
]

ABAS_SISTEMA = {'_Resumo', '_Pendentes', '_Observacoes', 'Sugestões Novas OS'}


def detectar_base():
    if os.environ.get('PCM_PROG_DIR') and os.path.isdir(os.environ['PCM_PROG_DIR']):
        return os.environ['PCM_PROG_DIR']
    try:
        sd = pathlib.Path(__file__).resolve().parent
        if (sd / 'BD_Relatório Semanal.xlsx').exists():
            return str(sd)
    except NameError:
        pass
    sub = pathlib.Path('4. O&M') / '11.Pré-Operação' / '6. PCM' / '09. Programação Semanal'
    home = pathlib.Path.home()
    for c in [
        home / 'GRID CO' / 'Grid Co. - Gridco' / sub,
        home / 'OneDrive - Grid Co' / sub,
        home / 'OneDrive - Gridco' / sub,
        home / 'OneDrive - GRID CO' / sub,
        home / 'OneDrive' / 'Grid Co. - Gridco' / sub,
        home / 'OneDrive' / sub,
    ]:
        try:
            if c.exists() and c.is_dir():
                return str(c)
        except OSError:
            pass
    return None


def carregar_lista_clusters(base):
    """Junta clusters do BD Semanal + AUXILIAR."""
    bd = os.path.join(base, 'BD_Relatório Semanal.xlsx')
    if not os.path.exists(bd):
        print(f'[ERRO] BD nao encontrado: {bd}')
        return []
    clusters = set()
    try:
        df_sem = pd.read_excel(bd, sheet_name='Semanal',
                                usecols=['Ativo Classificação 2'])
        for e in df_sem['Ativo Classificação 2'].dropna().unique():
            if isinstance(e, str) and e.strip():
                clusters.add(e.strip())
    except Exception as e:
        print(f'[AVISO] Falha lendo Semanal: {e}')
    try:
        df_aux = pd.read_excel(bd, sheet_name='AUXILIAR',
                                usecols=['Equipe Cluster'])
        for e in df_aux['Equipe Cluster'].dropna().unique():
            if isinstance(e, str) and e.strip():
                clusters.add(e.strip())
    except Exception as e:
        print(f'[AVISO] Falha lendo AUXILIAR: {e}')
    return sorted(clusters)


def listar_arquivos_semana(base, semanas_filtro=None):
    """Lista arquivos Programação Semana XX.xlsx. semanas_filtro = lista de ints."""
    out = []
    for f in os.listdir(base):
        if not f.lower().endswith('.xlsx'):
            continue
        if not f.lower().startswith('programa'):
            continue
        if 'semana' not in f.lower():
            continue
        m = re.search(r'semana\s*(\d+)', f.lower())
        if not m:
            continue
        num = int(m.group(1))
        if semanas_filtro and num not in semanas_filtro:
            continue
        out.append((num, os.path.join(base, f)))
    return sorted(out)


def migrar_arquivo(path, clusters_esperados, dry_run=False):
    """Adiciona abas faltantes. Retorna (n_existentes, n_criadas)."""
    nome = os.path.basename(path)
    print(f'\n=== {nome} ===')
    try:
        wb = load_workbook(path)
    except Exception as e:
        print(f'  [ERRO] Nao consegui abrir: {e}')
        return (0, 0)

    abas_atuais = set(wb.sheetnames)
    existentes = set()
    criadas = []

    for cluster in clusters_esperados:
        # Sheet name truncado conforme convencao do programacao_v7
        sheet_nm = re.sub(r'[\\/*?:\[\]]', '', cluster)[:31]
        if sheet_nm in abas_atuais:
            existentes.add(sheet_nm)
        else:
            criadas.append((cluster, sheet_nm))

    print(f'  Clusters esperados: {len(clusters_esperados)}')
    print(f'  Ja existentes:      {len(existentes)}')
    print(f'  A criar:            {len(criadas)}')

    if dry_run or not criadas:
        return (len(existentes), len(criadas))

    # Cria cada aba faltante com header padrao
    header_fill = PatternFill('solid', fgColor='4472C4')
    header_font = Font(bold=True, color='FFFFFF')
    align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for cluster, sheet_nm in criadas:
        ws = wb.create_sheet(sheet_nm)
        for j, col in enumerate(COLS_PADRAO_EQUIPE, start=1):
            c = ws.cell(row=1, column=j, value=col)
            c.fill = header_fill
            c.font = header_font
            c.alignment = align
        ws.freeze_panes = 'A2'
        print(f'    + criada: {sheet_nm}')

    # Reordena: abas _sistema primeiro, demais alfabeticamente
    ordem = []
    for a in ['_Resumo', '_Observacoes', '_Pendentes', 'Sugestões Novas OS']:
        if a in wb.sheetnames:
            ordem.append(a)
    outras = sorted([s for s in wb.sheetnames if s not in ordem])
    nova_ordem = ordem + outras
    # openpyxl reordena via _sheets
    wb._sheets = [wb[s] for s in nova_ordem]

    # Salva com retry
    for tent in range(3):
        try:
            wb.save(path)
            print(f'  [OK] Arquivo salvo. {len(criadas)} abas adicionadas.')
            return (len(existentes), len(criadas))
        except PermissionError as e:
            print(f'  [TENTATIVA {tent+1}/3] Arquivo bloqueado: {e}')
            import time
            time.sleep(3)
    print(f'  [ERRO] Nao consegui salvar (arquivo aberto no Excel?)')
    return (len(existentes), 0)


def main():
    print('=' * 60)
    print('Migracao: abas de cluster nos arquivos de Programacao')
    print('=' * 60)

    base = detectar_base()
    if not base:
        print('[ERRO] Pasta 09. Programacao Semanal nao encontrada.')
        sys.exit(1)
    print(f'Pasta base: {base}')

    # Filtra semanas via argv
    semanas_filtro = None
    if len(sys.argv) > 1:
        try:
            semanas_filtro = [int(x) for x in sys.argv[1:]]
            print(f'Filtro de semanas: {semanas_filtro}')
        except ValueError:
            print(f'[ERRO] Argumentos invalidos: {sys.argv[1:]}')
            sys.exit(1)

    clusters = carregar_lista_clusters(base)
    print(f'Clusters esperados (do BD): {len(clusters)}')
    if not clusters:
        print('[ERRO] Nenhum cluster encontrado no BD.')
        sys.exit(1)

    arquivos = listar_arquivos_semana(base, semanas_filtro)
    print(f'Arquivos a processar: {len(arquivos)}')
    if not arquivos:
        print('Nenhum arquivo correspondente.')
        sys.exit(0)

    total_criadas = 0
    for num, path in arquivos:
        _, criadas = migrar_arquivo(path, clusters)
        total_criadas += criadas

    print('\n' + '=' * 60)
    print(f'CONCLUIDO. {total_criadas} abas adicionadas em {len(arquivos)} arquivo(s).')
    print('=' * 60)


if __name__ == '__main__':
    main()
