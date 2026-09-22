# -*- coding: utf-8 -*-
"""
recalc_hh_ativo.py — Recalcula HH e MTTR sobre o TEMPO ATIVO (pausa descontada).

Lê o BD_Relatório Semanal.xlsx e usa, por OS, o campo 'MTTR (s)'
(= total_active_seconds da Fracttal, já sem pausa/verificação).
Onde 'MTTR (s)' estiver vazio, faz FALLBACK para 'Tempo de execução'
(o tempo aberto->fechado que usávamos) e marca essa OS como "sem tempo ativo".

Gera, para o cliente alvo (default Thopen) no período (default jan-jun/2026):
  - Cobertura: % de OS com tempo ativo real por natureza
  - Comparativo HH: aberto->fechado  x  tempo ativo  (quanto a pausa removeu)
  - MTTR médio por natureza (em tempo ativo)
  - % reativo (corretiva + religamento local) sobre o HH ativo
  - Abas Por Usina e Por Equipe
Saída: Comparativo_Tempo_Ativo_<cliente>_<periodo>.xlsx

Uso:
    python recalc_hh_ativo.py
    python recalc_hh_ativo.py --cliente Thopen --ini 2026-01-01 --fim 2026-07-01
"""
import os, sys, argparse, datetime as dt
from collections import defaultdict

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCRIPT_DIR)
import relatorio_clientes as R  # mesma pasta

def achar_pasta():
    for c in [os.environ.get("PCM_PROG_DIR"), SCRIPT_DIR, os.getcwd()]:
        if c and os.path.exists(os.path.join(c, "BD_Relatório Semanal.xlsx")):
            return c
    return SCRIPT_DIR

def num(v):
    try:
        f = float(v)
        return f
    except (TypeError, ValueError):
        return None

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cliente", default="Thopen")
    ap.add_argument("--ini", default="2026-01-01")
    ap.add_argument("--fim", default="2026-07-01")  # exclusivo
    ap.add_argument("--pasta", default=None)
    args = ap.parse_args()
    pasta = args.pasta or achar_pasta()
    ini = dt.datetime.fromisoformat(args.ini); fim = dt.datetime.fromisoformat(args.fim)
    cliente = args.cliente

    # força leitura fresca do BD (caso o cache esteja velho)
    cache = os.path.join(pasta, ".cache_bd.pkl")
    try:
        if os.path.exists(cache): os.remove(cache)
    except Exception:
        pass

    idx, rows, aux, cu = R.carregar(pasta)
    def g(r, k):
        i = idx.get(k); return r[i] if (i is not None and i < len(r)) else None
    tem_mttr = 'MTTR (s)' in idx
    tem_total = 'Tempo Total (s)' in idx
    tem_pausa = 'Tempo Pausado (s)' in idx

    def usina_de(r):
        c, u, cl, mob, cap = R.resolve_usina(g(r, 'Ativo Classificação 1'), aux)
        if u: return c, u
        loc = str(g(r, 'Localização ou parte de') or '')
        for p in [x.strip() for x in loc.split('/') if x.strip()]:
            if ' - ' in p:
                c2, u2, _, _, _ = R.resolve_usina(p, aux)
                if u2: return c2, u2
        return None, None

    def grupo(r):
        c = R.categoria(g(r, 'Tipo de tarefa') or '', g(r, 'Tarefa -> Classificação 2'))
        if c == 'Preventiva': return 'Preventiva'
        if c == 'Corretiva': return 'Corretiva'
        if c == 'Handover': return 'Handover'
        if c == 'Religamento': return 'Religamento local'
        if c == 'Religamento Remoto': return 'Religamento remoto'
        return 'Preditiva/Outras'

    def data_ref(r, grp):
        if grp == 'Handover':
            return (R.parse_dt(g(r, 'Data final')) or R.parse_dt(g(r, 'Data inicial'))
                    or R.parse_dt(g(r, 'Data Programada')))
        return R.mp_ref(dict(dfim=R.parse_dt(g(r, 'Data final')),
                             dprog=R.parse_dt(g(r, 'Data Programada')),
                             dini=R.parse_dt(g(r, 'Data inicial'))))

    GRUPOS = ['Corretiva', 'Religamento local', 'Preventiva', 'Handover',
              'Religamento remoto', 'Preditiva/Outras']
    # acumuladores
    cat = {k: {'n': 0, 'n_mttr': 0, 'h_aberto': 0.0, 'h_ativo': 0.0, 'h_pausa': 0.0,
               'mttr_soma': 0.0, 'mttr_n': 0} for k in GRUPOS}
    porusina = defaultdict(lambda: {'h_aberto': 0.0, 'h_ativo': 0.0, 'corr_ativo': 0.0,
                                    'reat_ativo': 0.0, 'n': 0, 'n_mttr': 0})
    porequipe = defaultdict(lambda: {'h_aberto': 0.0, 'h_ativo': 0.0, 'corr_ativo': 0.0,
                                     'reat_ativo': 0.0, 'usinas': set(), 'n': 0, 'n_mttr': 0})

    for r in rows:
        if str(g(r, 'Status')).strip() == 'Cancelado': continue
        cli, ufv = usina_de(r)
        if cli != cliente or not ufv: continue
        grp = grupo(r)
        d = data_ref(r, grp)
        if not (d and ini <= d < fim): continue

        # tempo aberto->fechado (referência atual)
        if tem_total and num(g(r, 'Tempo Total (s)')) and num(g(r, 'Tempo Total (s)')) > 0:
            h_aberto = num(g(r, 'Tempo Total (s)')) / 3600.0
        else:
            h_aberto = R.parse_dur(g(r, 'Tempo de execução'))
        # tempo ativo (MTTR real) com fallback
        mttr_s = num(g(r, 'MTTR (s)')) if tem_mttr else None
        tem_ativo = mttr_s is not None and mttr_s > 0
        h_ativo = (mttr_s / 3600.0) if tem_ativo else R.parse_dur(g(r, 'Tempo de execução'))
        # pausa
        if tem_pausa and num(g(r, 'Tempo Pausado (s)')) is not None:
            h_pausa = max(num(g(r, 'Tempo Pausado (s)')) / 3600.0, 0.0)
        else:
            h_pausa = max(h_aberto - h_ativo, 0.0)

        c = cat[grp]
        c['n'] += 1; c['h_aberto'] += h_aberto; c['h_ativo'] += h_ativo; c['h_pausa'] += h_pausa
        if tem_ativo:
            c['n_mttr'] += 1; c['mttr_soma'] += h_ativo; c['mttr_n'] += 1
        if grp == 'Religamento remoto':  # não consome HH de equipe
            continue
        pu = porusina[ufv]; pu['h_aberto'] += h_aberto; pu['h_ativo'] += h_ativo; pu['n'] += 1
        if tem_ativo: pu['n_mttr'] += 1
        if grp == 'Corretiva': pu['corr_ativo'] += h_ativo
        if grp in ('Corretiva', 'Religamento local'): pu['reat_ativo'] += h_ativo
        a2 = str(g(r, 'Ativo Classificação 2') or '').strip()
        if a2:
            pe = porequipe[a2]; pe['h_aberto'] += h_aberto; pe['h_ativo'] += h_ativo
            pe['usinas'].add(ufv.replace(cliente + ' - ', '')); pe['n'] += 1
            if tem_ativo: pe['n_mttr'] += 1
            if grp == 'Corretiva': pe['corr_ativo'] += h_ativo
            if grp in ('Corretiva', 'Religamento local'): pe['reat_ativo'] += h_ativo

    # ---------- console ----------
    total_n = sum(c['n'] for c in cat.values())
    total_mttr = sum(c['n_mttr'] for c in cat.values())
    print(f"\n== {cliente} | {args.ini} a {args.fim} ==")
    if not tem_mttr:
        print("AVISO: coluna 'MTTR (s)' não existe no BD — rode gerar_bd_via_api.py (com execuções). Usando fallback p/ tudo.")
    print(f"Cobertura de tempo ativo: {total_mttr}/{total_n} OS ({(100*total_mttr/total_n if total_n else 0):.0f}%)\n")
    hdr = f"{'Natureza':20} {'N':>5} {'cob%':>5} {'HH aberto':>10} {'HH ativo':>9} {'Pausa':>8} {'MTTR méd(h)':>11}"
    print(hdr)
    exec_aberto = exec_ativo = reat_aberto = reat_ativo = 0.0
    for k in GRUPOS:
        c = cat[k]
        cob = (100 * c['n_mttr'] / c['n']) if c['n'] else 0
        mttr_med = (c['mttr_soma'] / c['mttr_n']) if c['mttr_n'] else 0
        print(f"{k:20} {c['n']:>5} {cob:>4.0f}% {c['h_aberto']:>10.0f} {c['h_ativo']:>9.0f} {c['h_pausa']:>8.0f} {mttr_med:>11.2f}")
        if k != 'Religamento remoto':
            exec_aberto += c['h_aberto']; exec_ativo += c['h_ativo']
        if k in ('Corretiva', 'Religamento local'):
            reat_aberto += c['h_aberto']; reat_ativo += c['h_ativo']
    print("-" * len(hdr))
    print(f"{'HH de campo':20} {'':>5} {'':>5} {exec_aberto:>10.0f} {exec_ativo:>9.0f}")
    print(f"% Reativo  — aberto: {(100*reat_aberto/exec_aberto if exec_aberto else 0):.0f}%  |  ativo: {(100*reat_ativo/exec_ativo if exec_ativo else 0):.0f}%")
    print(f"% Corretiva— aberto: {(100*cat['Corretiva']['h_aberto']/exec_aberto if exec_aberto else 0):.0f}%  |  ativo: {(100*cat['Corretiva']['h_ativo']/exec_ativo if exec_ativo else 0):.0f}%")

    # ---------- xlsx ----------
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter as GL
        NAVY="191528"; LIME="A9DB21"; WHITE="FFFFFF"; INK="2A2533"; SEC="E7E9ED"; LTZ="F2F3F5"; AMBER="C0392B"
        thin=Side(style="thin",color="D9DBE0"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
        def st(c,b=False,s=10,col=INK,bg=None,al="center",nf=None):
            c.font=Font(name="Calibri",bold=b,size=s,color=col); c.alignment=Alignment(horizontal=al,vertical="center",wrap_text=(al=="left"))
            if bg:c.fill=PatternFill("solid",fgColor=bg)
            c.border=bd
            if nf:c.number_format=nf
        wb=Workbook(); ws=wb.active; ws.title="Resumo (tempo ativo)"; ws.sheet_view.showGridLines=False
        ws.merge_cells("A1:G1"); ws["A1"]=f"Tempo Ativo (pausa descontada) vs Aberto→Fechado — {cliente} — {args.ini[:7]} a {args.fim[:7]}"
        st(ws["A1"],True,12,WHITE,NAVY,"left"); ws.row_dimensions[1].height=22
        ws.merge_cells("A2:G2"); ws["A2"]=f"Tempo ativo = MTTR (s) da Fracttal (sem pausa/verificação). Onde vazio, usa 'Tempo de execução' (fallback). Cobertura: {total_mttr}/{total_n} OS ({(100*total_mttr/total_n if total_n else 0):.0f}%)."
        st(ws["A2"],False,8.5,"6B6878",LTZ,"left")
        cols=["Natureza","Nº OS","Cobertura ativo","HH aberto→fechado","HH tempo ativo","Pausa removida (h)","MTTR médio (h)"]
        for j,h in enumerate(cols): st(ws.cell(4,j+1,h),True,9.5,WHITE,NAVY,("left" if j==0 else "center"))
        rr=4
        for k in GRUPOS:
            c=cat[k]; rr+=1; bgc=WHITE if rr%2 else LTZ; ac=(k in('Corretiva','Religamento local'))
            st(ws.cell(rr,1,k),True if ac else False,9,(AMBER if ac else INK),al="left",bg=bgc)
            st(ws.cell(rr,2,c['n']),bg=bgc,nf="0")
            st(ws.cell(rr,3,(c['n_mttr']/c['n'] if c['n'] else 0)),bg=bgc,nf="0%")
            st(ws.cell(rr,4,round(c['h_aberto'])),bg=bgc,nf="0")
            st(ws.cell(rr,5,round(c['h_ativo'])),b=ac,col=(AMBER if ac else INK),bg=bgc,nf="0")
            st(ws.cell(rr,6,round(c['h_pausa'])),bg=bgc,nf="0")
            st(ws.cell(rr,7,((c['mttr_soma']/c['mttr_n']) if c['mttr_n'] else 0)),bg=bgc,nf="0.00")
        rr+=1
        st(ws.cell(rr,1,"HH de campo (exclui remoto)"),True,9,INK,al="left",bg=SEC)
        st(ws.cell(rr,2,""),bg=SEC); st(ws.cell(rr,3,""),bg=SEC)
        st(ws.cell(rr,4,round(exec_aberto)),True,9,INK,bg=SEC,nf="0")
        st(ws.cell(rr,5,round(exec_ativo)),True,9,INK,bg=SEC,nf="0")
        st(ws.cell(rr,6,round(exec_aberto-exec_ativo)),True,9,INK,bg=SEC,nf="0"); st(ws.cell(rr,7,""),bg=SEC)
        rr+=2
        st(ws.cell(rr,1,"% Reativo (Corretiva + Relig. local)"),True,9,INK,al="left")
        st(ws.cell(rr,4,(reat_aberto/exec_aberto if exec_aberto else 0)),bg=LTZ,nf="0.0%")
        st(ws.cell(rr,5,(reat_ativo/exec_ativo if exec_ativo else 0)),True,9,AMBER,bg=LTZ,nf="0.0%")
        rr+=1
        st(ws.cell(rr,1,"% Corretiva"),True,9,INK,al="left")
        st(ws.cell(rr,4,(cat['Corretiva']['h_aberto']/exec_aberto if exec_aberto else 0)),bg=LTZ,nf="0.0%")
        st(ws.cell(rr,5,(cat['Corretiva']['h_ativo']/exec_ativo if exec_ativo else 0)),True,9,AMBER,bg=LTZ,nf="0.0%")
        for j,w in enumerate([30,8,14,18,16,18,14]): ws.column_dimensions[GL(j+1)].width=w

        # Por Equipe
        w2=wb.create_sheet("Por Equipe"); w2.sheet_view.showGridLines=False
        w2.merge_cells("A1:H1"); w2["A1"]=f"Por Equipe — tempo ativo — {cliente}"; st(w2["A1"],True,11,WHITE,NAVY,"left")
        c2=["Equipe","Nº UFVs","UFVs","Nº OS","HH ativo","Corretiva ativo (h)","% Corretiva","% Reativo"]
        for j,h in enumerate(c2): st(w2.cell(3,j+1,h),True,9,WHITE,NAVY,("left" if j in(0,2) else "center"))
        rr=3
        for t in sorted(porequipe, key=lambda x:-porequipe[x]['h_ativo']):
            d=porequipe[t]; rr+=1; bgc=WHITE if rr%2 else LTZ
            st(w2.cell(rr,1,t),True,8.5,INK,al="left",bg=bgc); st(w2.cell(rr,2,len(d['usinas'])),bg=bgc,nf="0")
            st(w2.cell(rr,3,", ".join(sorted(d['usinas']))),s=8,al="left",bg=bgc)
            st(w2.cell(rr,4,d['n']),bg=bgc,nf="0"); st(w2.cell(rr,5,round(d['h_ativo'])),bg=bgc,nf="0")
            st(w2.cell(rr,6,round(d['corr_ativo'])),bg=bgc,nf="0")
            st(w2.cell(rr,7,(d['corr_ativo']/d['h_ativo'] if d['h_ativo'] else 0)),bg=bgc,nf="0.0%")
            st(w2.cell(rr,8,(d['reat_ativo']/d['h_ativo'] if d['h_ativo'] else 0)),bg=bgc,nf="0.0%")
        for j,w in enumerate([13,8,46,7,11,16,11,11]): w2.column_dimensions[GL(j+1)].width=w

        # Por Usina
        w3=wb.create_sheet("Por Usina"); w3.sheet_view.showGridLines=False
        w3.merge_cells("A1:G1"); w3["A1"]=f"Por Usina — tempo ativo — {cliente}"; st(w3["A1"],True,11,WHITE,NAVY,"left")
        c3=["Usina","Nº OS","Cobertura ativo","HH aberto→fechado","HH tempo ativo","Corretiva ativo (h)","% Corretiva"]
        for j,h in enumerate(c3): st(w3.cell(3,j+1,h),True,9,WHITE,NAVY,("left" if j==0 else "center"))
        rr=3
        for u in sorted(porusina, key=lambda x:-porusina[x]['corr_ativo']):
            d=porusina[u]; rr+=1; bgc=WHITE if rr%2 else LTZ
            st(w3.cell(rr,1,u.replace(cliente+' - ','')),True,8.5,INK,al="left",bg=bgc)
            st(w3.cell(rr,2,d['n']),bg=bgc,nf="0")
            st(w3.cell(rr,3,(d['n_mttr']/d['n'] if d['n'] else 0)),bg=bgc,nf="0%")
            st(w3.cell(rr,4,round(d['h_aberto'])),bg=bgc,nf="0")
            st(w3.cell(rr,5,round(d['h_ativo'])),bg=bgc,nf="0")
            st(w3.cell(rr,6,round(d['corr_ativo'])),bg=bgc,nf="0")
            st(w3.cell(rr,7,(d['corr_ativo']/d['h_ativo'] if d['h_ativo'] else 0)),bg=bgc,nf="0.0%")
        for j,w in enumerate([26,7,14,18,16,16,11]): w3.column_dimensions[GL(j+1)].width=w

        out=os.path.join(pasta, f"Comparativo_Tempo_Ativo_{cliente}_{args.ini[:7]}_{args.fim[:7]}.xlsx")
        wb.save(out); print(f"\nSALVO: {out}")
    except Exception as e:
        print(f"[xlsx] falhou: {e}")

if __name__ == "__main__":
    main()
