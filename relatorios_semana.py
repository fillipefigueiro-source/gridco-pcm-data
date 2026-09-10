# -*- coding: utf-8 -*-
"""
Relatórios de semana — usados pelo relatorio_diario.py (mesma engrenagem de PDF e e-mail).

  fechamento_semana(dados, week, gerado)  → "Fechamento da semana": aderência, temas críticos,
                                            reincidentes, religamentos, o que não coube, engenharia.
                                            Sexta-feira 17:00 BRT. 2 páginas.
  alerta_programacao(dados, week, gerado) → "Alerta da programação": 17 checagens sobre a semana
                                            recém-publicada (capacidade, dado ruim, decisões
                                            pendentes e viabilidade de campo). Dispara quando uma
                                            semana NOVA aparece no banco_dados.json.

Cada função devolve dict com: html, assunto, email (corpo curto em HTML), resumo (para o estado).
`dados` = {"bd", "gestao", "gerencial", "engenharia", "operacoes", "feriados"} já lidos.
Só regras fixas — nada de IA. Aprovado pelo Fillipe em 10/09/2026 (mockups S36 e S37).
"""
from __future__ import annotations

import collections
import html
import re
import unicodedata
from datetime import date, datetime, timedelta

META = 85.0
HH_DIA = 8.8
TIPOS_NAO_PROG = ("Corretiva", "Corretiva Emergencial", "Religamento", "Religamento Remoto")
DIAS_PT = ["Segunda-feira", "Terça-feira", "Quarta-feira", "Quinta-feira", "Sexta-feira", "Sábado", "Domingo"]
CAT = {"campo": "em campo", "remoto": "remoto (COS)", "queda": "queda de energia", "emergencial": "emergencial", "programado": "programado"}

h = lambda s: html.escape(str(s if s is not None else ""))
pct = lambda a, b: (100.0 * a / b) if b else 0.0
n = lambda v: f"{int(v):,}".replace(",", ".")
cor = lambda p: "ok" if p >= META else "mid" if p >= 70 else "mal"
cel = lambda p: f'<span class="cel {cor(p)}">{p:.0f}%</span>'
fin = lambda x: x.get("status") == "Finalizados"
ds = lambda s: date.fromisoformat(s[:10]) if s else None
usina_curta = lambda u: re.sub(r"^[^-]+ - ", "", u or "").rsplit(" - ", 1)[0]
tarefa_curta = lambda t: re.sub(r"^\[.*?\]\s*-?\s*", "", t or "")
janela_ok = lambda x: (x.get("h_fim") or "")[:2].isdigit() and int(x["h_fim"][:2]) < 24
hh = lambda x: (x.get("duracao") or 0) + (x.get("desloc") or 0)


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def _mins(s):
    try:
        a, b = str(s).split(":")
        return int(a) * 60 + int(b)
    except Exception:
        return None


def semana_key(d: date):
    return f"{d.isocalendar()[0]}-W{d.isocalendar()[1]:02d}"


def semana_seg(week: str):
    ano, num = week.split("-W")
    return date.fromisocalendar(int(ano), int(num), 1)


def tabela(linhas, cab, cols, vazio="nada encontrado"):
    if not linhas:
        return f'<div class="vazio">{vazio}</div>'
    return ('<table><thead><tr>' + ''.join(f'<th class="{c}">{t}</th>' for t, c in cab) + '</tr></thead><tbody>'
            + ''.join('<tr>' + ''.join(f'<td class="{c}">{v}</td>' for v, c in zip(l, cols)) + '</tr>' for l in linhas)
            + '</tbody></table>')


CSS_BASE = """
*{box-sizing:border-box} body{margin:0;font-family:"Segoe UI",Arial,sans-serif;font-size:9.5pt;color:#191528;line-height:1.3;background:#fff}
.pg{page-break-after:always} .pg:last-child{page-break-after:auto}
.top{display:flex;align-items:flex-start;gap:14px;padding-bottom:8px;margin-bottom:9px;border-bottom:3px solid #191528}
.top .logo{height:20px;width:auto;display:block;margin-top:4px} .logo-txt{font-weight:800;font-size:14pt} .top .t{flex:1}
.top h1{margin:0;font-size:15pt;font-weight:800;white-space:nowrap} .top .s{font-size:9pt;color:#68667d;white-space:nowrap}
.top .d{text-align:right;font-size:9pt;color:#3a3550;white-space:nowrap} .top .d b{display:block;font-size:12pt;color:#191528}
.kpis{display:grid;gap:6px;margin:8px 0 9px} .kpi{border:1px solid #e4e4ef;border-left:3px solid #cbcbdd;border-radius:6px;padding:5px 8px}
.kpi.g{border-left-color:#A9DB21} .kpi.r{border-left-color:#c2410c} .kpi.m{border-left-color:#a04408}
.kpi .v{font-size:14pt;font-weight:800;line-height:1.05} .kpi .n{font-size:7.8pt;color:#68667d;text-transform:uppercase;letter-spacing:.04em;margin-top:2px} .kpi .m{font-size:8pt;color:#68667d}
.sec{margin:8px 0 4px;font-size:8.5pt;font-weight:800;letter-spacing:.06em;text-transform:uppercase;color:#68667d;display:flex;align-items:center;gap:8px} .sec::after{content:"";flex:1;height:1px;background:#e4e4ef}
table{width:100%;border-collapse:collapse;font-size:8.4pt} td,th{white-space:nowrap;overflow:hidden;text-overflow:ellipsis;max-width:230px} td.quebra{white-space:normal;max-width:none}
th{text-align:left;font-size:7.4pt;text-transform:uppercase;letter-spacing:.04em;color:#68667d;padding:3px 6px;border-bottom:1px solid #cbcbdd;background:#f6f7fa}
td{padding:2.4px 6px;border-bottom:1px solid #eceef4;vertical-align:top} .r{text-align:right} th.r{text-align:right} .mono{font-family:Consolas,monospace;font-size:8.8pt}
.cel{display:inline-block;min-width:38px;text-align:center;padding:1px 5px;border-radius:4px;font-weight:700} .cel.ok{background:#e3f6ea;color:#1f7a4d} .cel.mid{background:#fdf0d5;color:#a04408} .cel.mal{background:#fde2e2;color:#b02525}
.g2{display:grid;grid-template-columns:1fr 1fr;gap:12px} .g3{display:grid;grid-template-columns:1.1fr 1fr 1fr;gap:12px}
.ins{margin:0;padding-left:16px} .ins li{margin:2.5px 0;font-size:9.1pt}
.bars{display:flex;gap:10px;align-items:flex-end;height:80px;padding:4px 6px 0} .bar{flex:1;display:flex;flex-direction:column;align-items:center;gap:3px;height:100%} .bar .col{flex:1;width:100%;display:flex;gap:2px;align-items:flex-end;justify-content:center}
.bar .col i,.bar .col b{display:block;width:14px;border-radius:2px 2px 0 0} .bar .col i{background:#cbcbdd} .bar .col b{background:#191528} .bar span{font-size:8.5pt;font-weight:700} .bar small{display:block;font-weight:400;color:#68667d;font-size:7.5pt}
.nota{background:#fbfbe8;border-left:3px solid #A9DB21;padding:5px 9px;font-size:8.6pt;color:#3a3550;margin-top:5px} .nota.aviso{background:#fdf0d5;border-left-color:#c2410c}
.pe{margin-top:9px;border-top:1px solid #e4e4ef;padding-top:6px;font-size:7.8pt;color:#68667d;display:flex;justify-content:space-between;gap:10px}
.al{color:#b02525;font-weight:700} .cinza{color:#68667d} .vazio{color:#68667d;font-size:8.6pt;padding:3px 2px}
.al-bloco{border:1px solid #e4e4ef;border-radius:6px;padding:5px 9px;margin:5px 0;break-inside:avoid}
.al-cab{display:flex;align-items:center;gap:8px;font-size:9.6pt} .al-txt{font-size:8.3pt;color:#3a3550;margin:1px 0 3px}
.sev{font-size:7.3pt;font-weight:800;text-transform:uppercase;letter-spacing:.05em;padding:1px 6px;border-radius:3px} .sev.mal{background:#fde2e2;color:#b02525} .sev.mid{background:#fdf0d5;color:#a04408} .sev.ok{background:#e3f6ea;color:#1f7a4d}
"""


def _pagina(titulo, css_extra, corpo, margem="11mm 12mm 12mm"):
    return (f'<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>{h(titulo)}</title>'
            f'<style>@page{{size:A4;margin:{margem}}}{CSS_BASE}{css_extra}</style></head><body>{corpo}</body></html>')


def _email(titulo, sub, kpis, ins, painel_url):
    celulas = "".join(f"<td style='padding:6px 10px;border:1px solid #e4e4ef;text-align:center'><div style='font-size:18px;font-weight:800'>{v}</div>"
                      f"<div style='font-size:11px;color:#68667d'>{r}</div></td>" for v, r in kpis)
    return (f"<div style='font-family:Segoe UI,Arial,sans-serif;font-size:14px;color:#191528'>"
            f"<h2 style='margin:0 0 4px'>{titulo}</h2><p style='margin:0 0 10px;color:#68667d'>{sub}</p>"
            f"<table style='border-collapse:collapse;margin-bottom:12px'><tr>{celulas}</tr></table>"
            f"<ul style='padding-left:18px'>{''.join(f'<li style=MG>{x}</li>' for x in ins)}</ul>"
            f"<p style='color:#68667d;font-size:12px;margin-top:18px'>Enviado pela plataforma PCM · <a href='{painel_url}'>abrir o painel</a>. "
            f"Responder este e-mail não registra nada.</p></div>").replace("style=MG", "style='margin:4px 0'")


# ═══════════════════════════════════════════════════════════════════════════
# FECHAMENTO DA SEMANA
# ═══════════════════════════════════════════════════════════════════════════
def fechamento_semana(dados, week, gerado: datetime, logo="", painel_url="https://pcm.gridco.com.br/novo.html"):
    bd, g, ger, eng = dados["bd"], dados["gestao"], dados["gerencial"], dados["engenharia"]
    W = {w["week"]: w for w in bd.get("semanas", [])}
    w = W.get(week)
    if not w:
        return None
    MON = semana_seg(week); SUN = MON + timedelta(6)
    plan = [x for x in w["rows"] if not x.get("foraDoPlano")]
    fora = [x for x in w["rows"] if x.get("foraDoPlano")]
    prog, feitas = len(plan), sum(map(fin, plan))
    ader = pct(feitas, prog)
    rol = sum(1 for x in plan if x.get("reprog") == "Sim")
    dias = [(d[:3], len([x for x in plan if x["dia"] == d]), sum(1 for x in plan if x["dia"] == d and fin(x))) for d in DIAS_PT[:5]]

    def bloco(k, top=8, minimo=1):
        c = collections.defaultdict(lambda: [0, 0, 0])
        for x in plan:
            kk = x.get(k) or "—"; c[kk][0] += 1; c[kk][1] += fin(x); c[kk][2] += (x.get("reprog") == "Sim")
        return [(kk, *v) for kk, v in sorted(c.items(), key=lambda kv: -kv[1][0]) if v[0] >= minimo][:top]
    cli = bloco("cliente"); resp = bloco("responsavel", 8)
    piores = sorted(bloco("cluster", 60, 10), key=lambda r: pct(r[2], r[1]))[:6]

    T = g.get("tarefas", [])
    tend = []
    for i in range(3, -1, -1):
        wk = semana_key(MON - timedelta(weeks=i))
        if wk not in W:
            continue
        mon = semana_seg(wk); sun = mon + timedelta(6)
        r = [x for x in W[wk]["rows"] if not x.get("foraDoPlano")]
        npw = sum(1 for t in T if t.get("tipo") in TIPOS_NAO_PROG and ds(t.get("criacao")) and mon <= ds(t["criacao"]) <= sun)
        rw = [e for e in ger.get("eventos", []) if ds(e.get("ini")) and mon <= ds(e["ini"]) <= sun]
        tend.append((wk[-2:], len(r), sum(map(fin, r)), sum(1 for x in r if x.get("reprog") == "Sim"), npw, len(rw),
                     sum((e.get("hDia") or 0) * (e.get("peso") or 1) for e in rw)))
    grp = collections.defaultdict(list)
    for x in plan:
        if (x.get("vezes") or 1) >= 4 and not fin(x):
            grp[x["os_id"]].append(x)
    reinc = sorted(grp.items(), key=lambda kv: -max(t.get("vezes") or 1 for t in kv[1]))
    npg = [t for t in T if t.get("tipo") in TIPOS_NAO_PROG and ds(t.get("criacao")) and MON <= ds(t["criacao"]) <= SUN]
    U = ger.get("usinas", [])
    usina_de = lambda e: U[e["u"]] if 0 <= e.get("u", -1) < len(U) else {}
    ev = [e for e in ger.get("eventos", []) if ds(e.get("ini")) and MON <= ds(e["ini"]) <= SUN]
    hrel = sum((e.get("hDia") or 0) * (e.get("peso") or 1) for e in ev)
    reg = collections.Counter(usina_de(e).get("regiao", "—") for e in ev); cat = collections.Counter(e.get("cat") for e in ev)
    hp = collections.Counter()
    for e in ev:
        hp[usina_de(e).get("usina", e.get("ativo", "—"))] += (e.get("hDia") or 0)
    causas = collections.Counter()
    for p in w.get("pendentes", []):
        m = re.sub(r"^\[\+\d+d em andamento\]\s*", "", p.get("motivo", ""))
        causas["MPA noturna sem dia definido" if "defina o dia" in m else "usina sem dia na semana" if "não recebeu dia" in m
               else "deslocada por OS grande" if "Deslocada por OS" in m else "rolagem sem capacidade" if "Rolagem" in m else "sem capacidade nos dias da usina"] += 1
    crit = [(a.get("nome"), a.get("usina"), a.get("nFalha")) for a in eng.get("ativos", []) if a.get("nivel") == "critico"]
    aten = sum(1 for a in eng.get("ativos", []) if a.get("nivel") == "atencao")
    kpi = eng.get("kpi") or {}

    t_ader = [pct(f, p) for _, p, f, *_ in tend]
    seguidas = 0
    for i in range(len(t_ader) - 1, 0, -1):
        if t_ader[i] < t_ader[i - 1]: seguidas += 1
        else: break
    tendencia = "caindo" if seguidas >= 2 else ("subindo" if len(t_ader) >= 3 and t_ader[-1] > t_ader[-2] > t_ader[-3] else "oscilando")
    pior_dia = min(dias, key=lambda d: pct(d[2], d[1])) if prog else ("—", 0, 0)
    resp_ok = [r for r in resp if r[0] not in ("—", "TBD")]
    ins = []
    if prog:
        seq = " → ".join(f"{x:.0f}%" for x in t_ader)
        ins.append(f"<b>Aderência de {ader:.0f}% na semana</b> — {n(feitas)} das {n(prog)} tarefas programadas finalizadas, "
                   f"{'acima' if ader >= META else 'abaixo'} da meta de {META:.0f}%. "
                   + (f"É a {['', 'segunda', 'terceira', 'quarta'][min(seguidas, 3)]} semana seguida em queda: {seq}." if tendencia == "caindo" else f"Últimas semanas: {seq}."))
        ins.append(f"<b>{n(rol)} tarefas rolaram</b> ({pct(rol, prog):.0f}% da programação). {n(len(reinc))} OS já rolaram 4 vezes ou mais e seguem abertas — "
                   f"são o tema crítico nº 1: cada uma dessas está há mais de um mês sendo empurrada.")
        if cli and resp_ok:
            ins.append(f"<b>{h(pior_dia[0])} foi o pior dia</b> ({pct(pior_dia[2], pior_dia[1]):.0f}% de {n(pior_dia[1])}). Por cliente, "
                       f"{h(min(cli, key=lambda r: pct(r[2], r[1]))[0])} teve a menor aderência; por supervisor, {h(min(resp_ok, key=lambda r: pct(r[2], r[1]))[0])}.")
    else:
        ins.append("<b>Semana sem programação carregada</b> no banco de dados — só o que entrou e os religamentos.")
    ant = tend[-2][4] if len(tend) >= 2 else None
    ins.append(f"<b>{n(len(npg))} OS não programadas entraram</b> ({' · '.join(f'{n(c)} {h(t)}' for t, c in collections.Counter(t['tipo'] for t in npg).most_common()) or 'nenhuma'})"
               + (f", {'mais' if len(npg) > ant else 'menos'} que na semana anterior ({n(ant)})." if ant is not None else ".") + " Isso consumiu HH que estava reservado ao plano.")
    if ev:
        top = hp.most_common(1)[0]
        ins.append(f"<b>{n(len(ev))} religamentos</b> somaram {hrel:.0f} h ponderadas de indisponibilidade. {h(top[0])} sozinha respondeu por {top[1]:.0f} h — vale abrir análise de causa raiz.")
    if fora:
        ins.append(f"<b>Fora do plano, a equipe executou {n(len(fora))} tarefas</b> que não estavam programadas. É trabalho real, mas invisível na aderência — e sinal de que o plano não reflete o que o campo faz.")
    if causas:
        ins.append(f"<b>{n(sum(causas.values()))} tarefas não couberam na semana</b>: {n(causas['MPA noturna sem dia definido'])} MPA noturnas esperam um dia nas observações e "
                   f"{n(causas['usina sem dia na semana'])} são de usinas que não receberam dia. Essas duas causas dependem do PCM, não do campo.")

    mx = max([p for _, p, _ in dias] or [1]) or 1
    barras = "".join(f'<div class="bar"><div class="col"><i style="height:{100*p/mx:.0f}%"></i><b style="height:{100*f/mx:.0f}%"></b></div><span>{d}<small>{n(f)}/{n(p)}</small></span></div>' for d, p, f in dias)
    label = w.get("label") or f"Semana {week[-2:]}"
    corpo = f"""
<div class="pg">
<div class="top">{logo}<div class="t"><h1>Fechamento da semana</h1><div class="s">Grid Co. · aderência, temas críticos e o que fica para a próxima semana</div></div>
  <div class="d"><b>{h(label)}</b>fechado {DIAS_PT[gerado.weekday()][:3].lower()} {gerado:%d/%m} às {gerado:%H:%M} · uso interno</div></div>
<div class="kpis" style="grid-template-columns:repeat(6,1fr)">
  <div class="kpi {'g' if ader >= META else 'r'}"><div class="v">{f'{ader:.0f}%' if prog else '—'}</div><div class="n">Aderência da semana</div><div class="m">{n(feitas)} de {n(prog)} · meta {META:.0f}%</div></div>
  <div class="kpi r"><div class="v">{n(rol)}</div><div class="n">Tarefas que rolaram</div><div class="m">{pct(rol, prog):.0f}% da programação</div></div>
  <div class="kpi"><div class="v">{n(len(fora))}</div><div class="n">Fora do plano</div><div class="m">executadas sem estar programadas</div></div>
  <div class="kpi r"><div class="v">{n(len(npg))}</div><div class="n">Não programadas</div><div class="m">corretivas e religamentos criados</div></div>
  <div class="kpi"><div class="v">{n(len(ev))}</div><div class="n">Religamentos</div><div class="m">{hrel:.0f} h ponderadas de parada</div></div>
  <div class="kpi r"><div class="v">{n(g.get('totalAtrasadas', 0))}</div><div class="n">Backlog atrasado</div><div class="m">de {n(g.get('totalAbertas', 0))} abertas no fechamento</div></div>
</div>
<div class="g2">
  <div><div class="sec">Aderência por dia — programadas × finalizadas</div><div class="bars">{barras}</div></div>
  <div><div class="sec">Quatro semanas — tendência</div>
    <table><thead><tr><th>Semana</th><th class="r">Prog.</th><th class="r">Aderência</th><th class="r">Rolaram</th><th class="r">Não prog.</th><th class="r">Relig.</th><th class="r">h parada</th></tr></thead><tbody>
    {''.join(f'<tr><td>S{s}</td><td class="r">{n(p)}</td><td class="r">{cel(pct(f,p))}</td><td class="r">{n(r)}</td><td class="r">{n(np_)}</td><td class="r">{n(rl)}</td><td class="r">{hh_:.0f}</td></tr>' for s,p,f,r,np_,rl,hh_ in tend)}
    </tbody></table></div>
</div>
<div class="g2">
  <div><div class="sec">Por cliente</div>{tabela([(h(k), n(p), n(f), n(r), cel(pct(f,p))) for k,p,f,r in cli], [("Cliente",""),("Prog.","r"),("Final.","r"),("Rolaram","r"),("Aderência","r")], ["","r","r","r","r"], "sem programação")}</div>
  <div><div class="sec">Por supervisor (Responsável O&amp;M)</div>{tabela([(h(k if k != "—" else "sem responsável no BD"), n(p), n(f), n(r), cel(pct(f,p))) for k,p,f,r in resp], [("Supervisor",""),("Prog.","r"),("Final.","r"),("Rolaram","r"),("Aderência","r")], ["","r","r","r","r"], "sem programação")}</div>
</div>
<div class="sec">Temas críticos da semana</div>
<ul class="ins">{''.join(f'<li>{x}</li>' for x in ins)}</ul>
<div class="pe"><span>Página 1 de 2 · Relatório gerado pela plataforma PCM · pcm.gridco.com.br</span><span>Aderência = finalizadas ÷ programadas (linhas do plano; fora do plano contadas à parte). Meta {META:.0f}%.</span></div>
</div>
<div class="pg">
<div class="top">{logo}<div class="t"><h1>Detalhe da semana</h1><div class="s">reincidentes, religamentos, o que não coube e a engenharia</div></div><div class="d"><b>{h(label)}</b>página 2 de 2</div></div>
<div class="sec">Reincidentes — OS com 4 rolagens ou mais que seguem abertas ({n(len(reinc))})</div>
{tabela([(f'#{h(k)}', h(usina_curta(L[0].get("usina"))[:24]), h(tarefa_curta(L[0].get("tarefa"))[:40]), h(L[0].get("tipo")), h(L[0].get("responsavel") or "—"), f'<span class="al">{max(t.get("vezes") or 1 for t in L)}ª</span>') for k,L in reinc[:10]],
        [("OS",""),("Usina",""),("Tarefa",""),("Tipo",""),("Responsável",""),("Rolagem","r")], ["mono","","","","","r"], "nenhuma OS com 4 rolagens ou mais em aberto")}
<div class="nota">Cada rolagem é uma semana perdida. Acima de 4, a causa raramente é capacidade: é peça, acesso, escopo ou a OS não cabe no dia. Vale decidir uma a uma na reunião de programação.</div>
<div class="g3">
  <div><div class="sec">Clusters com pior aderência (≥10 tarefas)</div>{tabela([(h(k), n(p), n(f), cel(pct(f,p))) for k,p,f,r in piores], [("Cluster",""),("Prog.","r"),("Final.","r"),("Ader.","r")], ["","r","r","r"])}</div>
  <div><div class="sec">Religamentos por região</div>{tabela([(h(k), n(v), f'{pct(v, len(ev)):.0f}%') for k,v in reg.most_common()], [("Região",""),("Eventos","r"),("Parte","r")], ["","r","r"], "nenhum religamento")}
    {f'<div class="nota">{" · ".join(f"<b>{n(v)}</b> {CAT.get(k,k)}" for k,v in cat.most_common())}</div>' if cat else ''}</div>
  <div><div class="sec">Usinas com mais horas de parada</div>{tabela([(h(str(k)[:24]), f'{v:.0f} h') for k,v in hp.most_common(6)], [("Usina",""),("Diurnas","r")], ["","r"], "—")}</div>
</div>
<div class="g2">
  <div><div class="sec">Não programadas da semana — por cliente</div>{tabela([(h(k or "—"), n(v), f'{pct(v, len(npg)):.0f}%') for k,v in collections.Counter(t.get("cliente") for t in npg).most_common(6)], [("Cliente",""),("OS","r"),("Parte","r")], ["","r","r"], "nenhuma")}</div>
  <div><div class="sec">O que não coube na semana — por causa</div>{tabela([(h(k), n(v)) for k,v in causas.most_common()], [("Causa",""),("Tarefas","r")], ["","r"], "tudo coube")}
    <div class="nota">"MPA noturna sem dia" e "usina sem dia" dependem do PCM: defina o dia nas Observações da Semana e a próxima programação absorve.</div></div>
</div>
<div class="sec">Engenharia — Confiabilidade no fechamento</div>
<ul class="ins">
  <li><b>{n(kpi.get('ativosSinal', 0))} ativos com sinal de recorrência</b>, {n(len(crit))} crítico(s) e {n(aten)} em atenção. Crítico: {'; '.join(f'{h(a)} ({h(u)}, {f} falhas em 30 dias)' for a,u,f in crit) or 'nenhum'}.</li>
  <li><b>MTTR de {kpi.get('mttr', 0):.1f} h</b> e MTBF de {n(kpi.get('mtbf', 0))} h na frota monitorada. {n(kpi.get('abaixo85', 0))} ativos abaixo de 85% de disponibilidade.</li>
</ul>
<div class="pe"><span>Página 2 de 2 · o cliente não recebe este relatório — ele vê o resultado, não o processo</span><span>Fontes: Programação Semanal, Gestão PCM, Gerencial e Confiabilidade (Fracttal via robô)</span></div>
</div>"""
    kpis = [(f"{ader:.0f}%" if prog else "—", f"aderência ({n(feitas)} de {n(prog)})"), (n(rol), "rolaram"), (n(len(reinc)), "OS com 4+ rolagens"),
            (n(len(npg)), "não programadas"), (n(len(ev)), "religamentos"), (n(len(fora)), "fora do plano")]
    return {
        "html": _pagina(f"Fechamento da semana · {label}", "", corpo),
        "assunto": f"Fechamento da semana · {label} · aderência {ader:.0f}% · {n(rol)} rolaram · {n(len(reinc))} OS reincidentes",
        "email": _email(f"Fechamento da semana — {label}", "O relatório completo vai no PDF em anexo (2 páginas).", kpis, ins, painel_url),
        "resumo": {"prog": prog, "fin": feitas, "aderencia": round(ader, 1), "rolaram": rol, "reincidentes": len(reinc), "naoProg": len(npg), "relig": len(ev), "foraPlano": len(fora)},
    }


# ═══════════════════════════════════════════════════════════════════════════
# ALERTA DA PROGRAMAÇÃO
# ═══════════════════════════════════════════════════════════════════════════
def alerta_programacao(dados, week, gerado: datetime, logo="", painel_url="https://pcm.gridco.com.br/novo.html"):
    bd, op, feriados = dados["bd"], dados.get("operacoes") or {}, dados.get("feriados") or []
    W = {w["week"]: w for w in bd.get("semanas", [])}
    w = W.get(week)
    if not w:
        return None
    MON = semana_seg(week)
    rows = [x for x in w["rows"] if not x.get("foraDoPlano")]
    resumo = w.get("resumo") or {}
    P = w.get("pendentes", [])
    A = []

    # ── dado e capacidade ──
    imp = collections.defaultdict(list)
    for x in rows:
        if (x.get("duracao") or 0) > 8 or not janela_ok(x):
            imp[x["os_id"]].append(x)
    A.append(("alta", f"Duração implausível ou janela impossível — {n(len(imp))} OS, {n(sum(len(L) for L in imp.values()))} tarefas",
              "Tarefa com mais de 8 h numa janela, ou fim depois da meia-noite. Distorce o HH do cluster e empurra as outras. Corrigir a duração no plano de manutenção do Fracttal.",
              [(f'#{h(k)}', h(usina_curta(L[0].get("usina"))[:22]), h(tarefa_curta(L[0].get("tarefa"))[:30]), f'{L[0].get("duracao") or 0:.1f} h', h(L[0].get("h_ini")), h(L[0].get("h_fim")), n(len(L)))
               for k, L in sorted(imp.items(), key=lambda kv: -(kv[1][0].get("duracao") or 0))[:4]],
              [("OS", ""), ("Usina", ""), ("Tarefa", ""), ("Duração", "r"), ("Início", ""), ("Fim", ""), ("Linhas", "r")], ["mono", "", "", "r", "mono", "mono", "r"]))
    cap = sorted([(k, v["hh_util"], v["hh_disp"], v.get("pend", 0)) for k, v in resumo.items() if v.get("hh_disp") and v["hh_util"] > v["hh_disp"]], key=lambda t: -t[1] / max(t[2], 1))
    A.append(("alta", f"Clusters acima da capacidade semanal — {n(len(cap))} de {n(len(resumo))}",
              "HH programado maior que as horas da semana do cluster. Vai rolar por construção. Ou tira tarefa, ou muda a capacidade no BD de Operações.",
              [(h(k), f"{u:.0f} h", f"{d} h", f'<span class="cel {"mal" if u/d > 1.5 else "mid"}">{pct(u, d):.0f}%</span>', n(p)) for k, u, d, p in cap[:5]],
              [("Cluster", ""), ("Programado", "r"), ("Capacidade", "r"), ("Uso", "r"), ("Ficou fora", "r")], ["", "r", "r", "r", "r"]))
    dc = collections.defaultdict(float)
    for x in rows:
        if janela_ok(x):
            dc[(x["dia"], x.get("cluster") or "—")] += hh(x)
    over = sorted([(k, v) for k, v in dc.items() if v > HH_DIA], key=lambda kv: -kv[1])
    A.append(("media", f"Dias com mais de {HH_DIA:.1f} h num cluster — {n(len(over))} de {n(len(dc))} combinações dia × cluster",
              "Mais horas num dia do que uma pessoa executa. O programador aceita porque a capacidade é semanal; o campo não. Redistribuir entre os dias.",
              [(h(k[0][:3]), h(k[1]), f"{v:.0f} h", f'<span class="cel {"mal" if v > 2*HH_DIA else "mid"}">{v/HH_DIA:.1f}×</span>') for k, v in over[:5]],
              [("Dia", ""), ("Cluster", ""), ("Programado", "r"), ("Vs. 1 pessoa", "r")], ["", "", "r", "r"]))
    pd_ = collections.Counter(x["dia"] for x in rows)
    dias = [(d[:3], pd_.get(d, 0)) for d in DIAS_PT[:6] if pd_.get(d, 0)]
    if dias:
        mx, mn = max(v for _, v in dias), min(v for _, v in dias)
        A.append(("media" if pct(mx, len(rows)) >= 35 else "baixa", f"Semana desbalanceada — {h(max(dias, key=lambda t: t[1])[0])} tem {n(mx)} tarefas, {h(min(dias, key=lambda t: t[1])[0])} tem {n(mn)}",
                  f"{pct(mx, len(rows)):.0f}% da semana num só dia. Um dia de chuva ou uma corretiva grande derruba a aderência da semana inteira.",
                  [tuple(f"<b>{h(d)}</b> {n(v)}" for d, v in dias)], [(d, "") for d, _ in dias], [""] * len(dias)))
    byk = collections.defaultdict(list)
    for x in rows:
        if janela_ok(x) and x.get("h_ini") and x.get("h_fim"):
            byk[(x["dia"], x.get("cluster"))].append((_mins(x["h_ini"]), _mins(x["h_fim"]), x["os_id"], x.get("usina")))
    sob = []
    for k, L in byk.items():
        L.sort(key=lambda t: (t[0] or 0, t[1] or 0))
        for i in range(1, len(L)):
            if L[i][0] is not None and L[i - 1][1] is not None and L[i][0] < L[i - 1][1] and L[i][2] != L[i - 1][2]:
                sob.append((k[0][:3], k[1], L[i - 1][2], L[i][2], L[i - 1][3] != L[i][3]))
    A.append(("media", f"Janelas sobrepostas no mesmo cluster e dia — {n(len(sob))} pares",
              "Duas OS diferentes ocupando a mesma hora da mesma equipe. Se são usinas diferentes, uma delas vai rolar.",
              [(h(d), h(c), f'#{h(a)} × #{h(b)}', "usinas diferentes" if dif else "mesma usina") for d, c, a, b, dif in sob[:4]],
              [("Dia", ""), ("Cluster", ""), ("OS", ""), ("Onde", "")], ["", "", "mono", ""]))
    v5 = collections.defaultdict(list)
    for x in rows:
        if (x.get("vezes") or 1) >= 5:
            v5[x["os_id"]].append(x)
    A.append(("media", f"OS com 5 rolagens ou mais programadas de novo — {n(len(v5))} OS",
              "Já rolaram cinco vezes ou mais e voltaram para a semana sem nada mudar. A chance de rolar de novo é alta. Decidir: peça, acesso, escopo ou tirar do plano.",
              [(f'#{h(k)}', h(usina_curta(L[0].get("usina"))[:22]), h(tarefa_curta(L[0].get("tarefa"))[:34]), h(L[0].get("responsavel") or "—"), f'<b class="al">{max(t.get("vezes") or 1 for t in L)}ª</b>')
               for k, L in sorted(v5.items(), key=lambda kv: -max(t.get("vezes") or 1 for t in kv[1]))[:5]],
              [("OS", ""), ("Usina", ""), ("Tarefa", ""), ("Responsável", ""), ("Rolagem", "r")], ["mono", "", "", "", "r"]))
    tbd = [x for x in rows if (x.get("responsavel") or "TBD") == "TBD"]
    A.append(("baixa", f"Tarefas sem responsável definido — {n(len(tbd))} (TBD)",
              "O cluster não tem Responsável O&M no BD de Operações. Ninguém recebe o alerta do dia dessas tarefas.",
              [(h(c or "—"), n(v)) for c, v in collections.Counter(x.get("cluster") for x in tbd).most_common(6)], [("Cluster", ""), ("Tarefas", "r")], ["", "r"]))
    mpa = collections.defaultdict(list); semdia = collections.Counter(); desloc = collections.Counter()
    for p in P:
        m = re.sub(r"^\[\+\d+d em andamento\]\s*", "", p.get("motivo", ""))
        if "defina o dia" in m:
            mpa[p["os_id"]].append(p)
        elif "não recebeu dia" in m:
            mm = re.search(r'Usina "([^"]+)"', m); semdia[mm.group(1) if mm else "?"] += 1
        elif "Deslocada por OS" in m:
            mm = re.search(r"OS #(\d+)", m); desloc[mm.group(1) if mm else "?"] += 1
    A.append(("alta" if (mpa or semdia) else "baixa", f"Ficou fora da semana por decisão pendente — {n(len(mpa))} OS de MPA noturna e {n(len(semdia))} usinas sem dia",
              "O programador não decide sozinho: MPA noturna precisa de um dia nas Observações da Semana; usina sem dia precisa entrar na distribuição do cluster. Enquanto isso, essas tarefas envelhecem no backlog.",
              [x for x in [("MPA noturna sem dia", ", ".join(f"#{k}" for k in sorted(mpa)[:10]) + (f" +{len(mpa)-10}" if len(mpa) > 10 else "")) if mpa else None,
                           ("Usinas sem dia na semana", ", ".join(usina_curta(u) for u, _ in semdia.most_common(6)) + (f" +{len(semdia)-6}" if len(semdia) > 6 else "")) if semdia else None,
                           ("Deslocadas por OS grande", " · ".join(f"#{k} empurrou {n(v)}" for k, v in desloc.most_common(4))) if desloc else None] if x],
              [("Situação", ""), ("Quem", "")], ["", "quebra"]))
    ufs = collections.Counter((x.get("usina") or "").rsplit(" - ", 1)[-1].strip()[:2] for x in rows)
    fl = []
    for t, uf, d, nome in feriados:
        if not (MON <= d <= MON + timedelta(5)):
            continue
        prog_dia = sum(1 for x in rows if x["dia"] == DIAS_PT[d.weekday()] and (uf == "TODOS" or (x.get("usina") or "").rsplit(" - ", 1)[-1].strip()[:2] == uf))
        fl.append((h(DIAS_PT[d.weekday()][:3] + d.strftime(" %d/%m")), h(nome), h(uf), n(prog_dia), prog_dia))
    A.append(("alta" if any(l[4] for l in fl) else "baixa", f"Feriados na semana — {n(len(fl))}",
              "Tarefa programada em feriado da UF da usina só roda se a equipe trabalhar no feriado. Confirmar ou mover.",
              [l[:4] for l in fl], [("Data", ""), ("Feriado", ""), ("UF", ""), ("Tarefas no dia", "r")], ["", "", "", "r"], ))
    qa = [q for q in w.get("qualidade", []) if q.get("tipo") != "REMOVIDA"]
    A.append(("baixa", f"Avisos de cadastro — {n(len(qa))} · {n(len(w.get('qualidade', [])) - len(qa))} linhas de teste removidas",
              "Erros de nome no Fracttal que fazem a distribuição de dias contar errado.",
              [(h(q.get("item", "")[:34]), h(q.get("detalhe", "")[:60]), h(q.get("acao", "")[:40])) for q in qa], [("Item", ""), ("Problema", ""), ("Ação", "")], ["", "", ""]))

    # ── viabilidade de campo ──
    cid = {_norm(u_.get("usina")): (u_.get("cidade") or "").strip() for u_ in op.get("usinas", [])}
    def cidade(us):
        c = cid.get(_norm(usina_curta(us)))
        if not c:  # recuo: cidade pelo nome da usina sem o sufixo numérico, senão a UF
            base = re.sub(r"\s*\d+\s*(\(.*\))?$", "", usina_curta(us))
            c = cid.get(_norm(base)) or ("UF " + (us or "").rsplit(" - ", 1)[-1].strip()[:2])
        return c
    gd = collections.defaultdict(list)
    for x in rows:
        if janela_ok(x):
            gd[(x["dia"], x.get("cluster") or "—")].append(x)
    multi = []
    for k, L in gd.items():
        cids = collections.Counter(cidade(t.get("usina")) for t in L)
        if len(cids) >= 2:
            multi.append((k, cids, sum(t.get("desloc") or 0 for t in L), len(L)))
    multi.sort(key=lambda t: (-len(t[1]), -t[2]))
    A.append(("alta" if any(len(m[1]) >= 3 for m in multi) else "media" if multi else "baixa",
              f"Equipe em mais de uma cidade no mesmo dia — {n(len(multi))} de {n(len(gd))} dias × cluster ({n(sum(1 for m in multi if len(m[1]) >= 3))} com 3 ou mais)",
              "A mesma equipe programada em usinas de cidades diferentes no mesmo dia. Duas pode ser rotina; três ou mais raramente cabe com deslocamento, almoço e apontamento.",
              [(h(k[0][:3]), h(k[1]), n(len(c)), h(" · ".join(f"{ci} ({v})" for ci, v in c.most_common(4))[:58]), f"{d:.1f} h", n(nt)) for k, c, d, nt in multi[:6]],
              [("Dia", ""), ("Cluster", ""), ("Cidades", "r"), ("Onde (tarefas)", ""), ("Desloc.", "r"), ("Tarefas", "r")], ["", "", "r", "", "r", "r"]))
    dsl = sorted([(k, sum(t.get("desloc") or 0 for t in L), len({cidade(t.get("usina")) for t in L})) for k, L in gd.items() if sum(t.get("desloc") or 0 for t in L) > 3], key=lambda t: -t[1])
    A.append(("media" if dsl else "baixa", f"Mais de 3 h de deslocamento num dia — {n(len(dsl))} dias × cluster",
              "Deslocamento somado das tarefas do dia. Acima de 3 h, sobra menos de 6 h de execução real.",
              [(h(k[0][:3]), h(k[1]), f"{d:.1f} h", n(c)) for k, d, c in dsl[:5]], [("Dia", ""), ("Cluster", ""), ("Deslocamento", "r"), ("Cidades", "r")], ["", "", "r", "r"]))
    seq = []
    for k, L in gd.items():
        Ls = sorted(L, key=lambda t: _mins(t.get("h_ini")) or 0)
        for i in range(1, len(Ls)):
            a_, b_ = Ls[i - 1], Ls[i]
            if a_.get("usina") != b_.get("usina") and _mins(b_.get("h_ini")) is not None and _mins(a_.get("h_fim")) is not None:
                gap = _mins(b_["h_ini"]) - _mins(a_["h_fim"]); need = (b_.get("desloc") or 0) * 60
                if gap < 0 or gap < need - 5:
                    seq.append((k[0][:3], k[1], a_["os_id"], b_["os_id"], gap, need, cidade(a_.get("usina")), cidade(b_.get("usina"))))
    A.append(("media" if seq else "baixa", f"Troca de usina sem tempo de deslocamento — {n(len(seq))} casos",
              "A tarefa seguinte começa em outra usina antes de a anterior terminar, ou sem o tempo de estrada que o próprio plano calcula.",
              [(h(d), h(c), f"#{h(a)} → #{h(b)}", h(f"{ca} → {cb}"[:30]), f"{g_} min", f"{nd:.0f} min") for d, c, a, b, g_, nd, ca, cb in seq[:5]],
              [("Dia", ""), ("Cluster", ""), ("OS", ""), ("Trajeto", ""), ("Folga", "r"), ("Precisa", "r")], ["", "", "mono", "", "r", "r"]))
    fh = collections.defaultdict(list)
    for x in rows:
        if janela_ok(x) and "NOTURNO" not in x["dia"].upper() and (x.get("tipo") or "") != "MPA" and _mins(x.get("h_ini")) is not None \
                and (_mins(x["h_ini"]) < 6 * 60 or (_mins(x.get("h_fim")) or 0) > 19 * 60):
            fh[x["os_id"]].append(x)
    A.append(("media" if fh else "baixa", f"Fora do horário de campo (06:00–19:00) — {n(len(fh))} OS, {n(sum(len(L) for L in fh.values()))} tarefas",
              "Diurna programada antes das 6 h ou terminando depois das 19 h. MPA noturna e janelas marcadas NOTURNO ficam fora desta checagem.",
              [(f'#{h(k)}', h(usina_curta(L[0].get("usina"))[:22]), h(tarefa_curta(L[0].get("tarefa"))[:30]), h(L[0].get("tipo")), h(L[0].get("h_ini")), h(L[0].get("h_fim"))) for k, L in list(fh.items())[:5]],
              [("OS", ""), ("Usina", ""), ("Tarefa", ""), ("Tipo", ""), ("Início", ""), ("Fim", "")], ["mono", "", "", "", "mono", "mono"]))
    cnt = sorted([(k, len(L), len({t["os_id"] for t in L})) for k, L in gd.items() if len(L) > 12], key=lambda t: -t[1])
    A.append(("media" if cnt else "baixa", f"Mais de 12 tarefas para uma equipe num dia — {n(len(cnt))} dias × cluster",
              "Mesmo que o HH feche, cada tarefa tem abertura, execução, foto e fechamento no Fracttal. Acima de 12 por dia o apontamento não acompanha.",
              [(h(k[0][:3]), h(k[1]), n(c), n(o_)) for k, c, o_ in cnt[:5]], [("Dia", ""), ("Cluster", ""), ("Tarefas", "r"), ("OS", "r")], ["", "", "r", "r"]))
    fds = [x for x in rows if x["dia"].startswith(("Sáb", "Dom"))]
    A.append(("baixa", f"Programação em fim de semana — {n(len(fds))} tarefas", "Só vale se a equipe trabalha no sábado. Senão, rola por construção.",
              [(h(c or "—"), n(v)) for c, v in collections.Counter(x.get("cluster") for x in fds).most_common(4)], [("Cluster", ""), ("Tarefas", "r")], ["", "r"]))
    pes = {c.get("cluster"): c.get("pessoas", 0) for c in op.get("clusters", [])}
    sp = collections.Counter(x.get("cluster") for x in rows if pes.get(x.get("cluster"), 1) == 0)
    A.append(("alta" if sp else "baixa", f"Cluster com programação mas sem pessoa no BD de Operações — {n(len(sp))}",
              "A capacidade da semana vem de zero pessoas. Ou o colaborador não está na Relação Geral, ou o cluster está com o nome diferente.",
              [(h(c), n(v)) for c, v in sp.most_common(5)], [("Cluster", ""), ("Tarefas", "r")], ["", "r"]))

    ORD = {"alta": 0, "media": 1, "baixa": 2}
    SEV = {"alta": ("Alta", "mal"), "media": ("Média", "mid"), "baixa": ("Baixa", "ok")}
    A.sort(key=lambda a: ORD[a[0]])
    cont = collections.Counter(a[0] for a in A)
    hh_tot = sum(hh(x) for x in rows if janela_ok(x)); cap_tot = sum(v.get("hh_disp", 0) for v in resumo.values())
    blocos = "".join(f'<div class="al-bloco"><div class="al-cab"><span class="sev {SEV[s][1]}">{SEV[s][0]}</span><b>{t}</b></div>'
                     f'<div class="al-txt">{r}</div>{tabela(l, cab, cols)}</div>' for s, t, r, l, cab, cols in A)
    label = w.get("label") or f"Semana {week[-2:]}"
    corpo = f"""
<div class="top" style="border-bottom-color:#a04408">{logo}<div class="t"><h1>Alerta da programação</h1><div class="s">Grid Co. · checagem da programação recém-gerada, antes de ir para o campo</div></div>
  <div class="d"><b>{h(label)}</b>gerado {DIAS_PT[gerado.weekday()][:3].lower()} {gerado:%d/%m} às {gerado:%H:%M} · uso interno (PCM)</div></div>
<div class="kpis" style="grid-template-columns:repeat(5,1fr)">
  <div class="kpi"><div class="v">{n(len(rows))}</div><div class="n">Tarefas programadas</div><div class="m">em {n(len({x['os_id'] for x in rows}))} OS · {n(len(resumo))} clusters</div></div>
  <div class="kpi {'r' if hh_tot > cap_tot else 'g'}"><div class="v">{pct(hh_tot, cap_tot):.0f}%</div><div class="n">HH ÷ capacidade</div><div class="m">{hh_tot:.0f} h de {n(cap_tot)} h na semana</div></div>
  <div class="kpi r"><div class="v">{n(cont['alta'])}</div><div class="n">Severidade alta</div><div class="m">resolver antes de segunda</div></div>
  <div class="kpi m"><div class="v">{n(cont['media'])}</div><div class="n">Severidade média</div><div class="m">vão virar rolagem</div></div>
  <div class="kpi m"><div class="v">{n(len(multi))}</div><div class="n">Dias em 2+ cidades</div><div class="m">mesma equipe · {n(len(P))} tarefas não couberam</div></div>
</div>
{blocos}
<div class="pe"><span>Relatório gerado pela plataforma PCM · pcm.gridco.com.br · o cliente não recebe este relatório</span><span>Fonte: Programação Semanal publicada (banco_dados.json) e BD de Operações. Corrija na planilha ou nas Observações e rode o programador de novo.</span></div>"""
    altas = [t for s, t, *_ in A if s == "alta"]
    kpis = [(n(len(rows)), "tarefas"), (f"{pct(hh_tot, cap_tot):.0f}%", "HH ÷ capacidade"), (n(cont["alta"]), "alertas altos"), (n(cont["media"]), "médios"), (n(len(multi)), "dias em 2+ cidades")]
    return {
        "html": _pagina(f"Alerta da programação · {label}", "", corpo, "10mm 11mm 10mm"),
        "assunto": f"Alerta da programação · {label} · {n(cont['alta'])} altos, {n(cont['media'])} médios · HH {pct(hh_tot, cap_tot):.0f}% da capacidade",
        "email": _email(f"Alerta da programação — {label}", "As checagens completas vão no PDF em anexo. Corrija na planilha ou nas Observações e rode o programador de novo.", kpis, [f"<b>Alta:</b> {t}" for t in altas], painel_url),
        "resumo": {"tarefas": len(rows), "hhPct": round(pct(hh_tot, cap_tot)), "alta": cont["alta"], "media": cont["media"], "baixa": cont["baixa"], "multiCidade": len(multi)},
    }


def ler_feriados(caminho):
    """[(tipo, uf, date, nome)] da planilha de feriados. Sem planilha ou sem openpyxl → []."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(caminho, read_only=True)
        ws = wb[wb.sheetnames[0]]
        out = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i < 2:
                continue
            for off in (0, 6):
                if len(r) > off + 4 and isinstance(r[off + 2], datetime):
                    out.append((r[off], str(r[off + 1] or "").strip(), r[off + 2].date(), r[off + 4]))
        return out
    except Exception:
        return []
