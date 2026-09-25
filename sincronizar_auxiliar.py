# -*- coding: utf-8 -*-
"""
sincronizar_auxiliar.py — leva a AUXILIAR VIVA para o repositório, sem dado sensível.

Os geradores do painel rodam no GitHub Actions e não enxergam o OneDrive: leem
"AUXILIAR - FABRICIO.xlsx" e "AUXILIAR_gestao.xlsx" de dentro do repositório,
que estavam congeladas desde julho/agosto (supervisores desatualizados → TBD).
Este script, rodado na máquina do PCM junto com o publicar_mpas (15/15 min):
  1. lê a AUXILIAR viva do atalho do OneDrive (mesmo padrão da Gerencial);
  2. REMOVE as colunas sensíveis (CNPJ, receita, contatos, acessos, endereço) —
     o repositório é PÚBLICO;
  3. grava as duas cópias do repo com a mesma estrutura (aba Operacoes_1,
     cabeçalho por nome) só quando o conteúdo mudou.
Quem publica é o chamador (sync_repo). Uso manual:  py -3 sincronizar_auxiliar.py
"""
import hashlib
import json
import os
import sys

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
CANDIDATOS = [
    os.environ.get("AUXILIAR_VIVA_PATH") or "",
    os.path.join(os.path.expanduser("~"), "OneDrive - GRID CO", "Shortcuts", "GRID CO_ - 4. O&M",
                 "11.Pré-Operação", "6. PCM", "09. Programação Semanal", "AUXILIAR - FABRICIO.xlsx"),
]
DESTINOS = [os.path.join(AQUI, "AUXILIAR - FABRICIO.xlsx"), os.path.join(AQUI, "AUXILIAR_gestao.xlsx")]
MARCA = os.path.join(AQUI, "_auxiliar_hash.txt")

# colunas que NUNCA vão para o repositório público (comparação sem acento/caixa)
SENSIVEIS_EXATAS = {
    "contrato assinado", "razao social", "cnpj de faturamento", "prazo contratual (meses)",
    "receita mensal", "receita contratual", "cep", "endereco", "nome spe", "instalacao",
    "conta contrato", "cnpj", "ucs", "maps", "databook", "personalizar",
    "empresa seguranca local", "empresa seguranca remota", "prestador de servico - internet",
    "acesso para comando remoto", "acesso para visualizacao (site de monitoramento)",
}
SENSIVEIS_PREFIXO = ("contato", "url imagem")


def _sem_acento(s):
    import unicodedata
    return unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().strip().lower()


def _sensivel(nome):
    n = " ".join(_sem_acento(nome).split())
    return n in SENSIVEIS_EXATAS or n.startswith(SENSIVEIS_PREFIXO)


def log(m):
    print(f"[auxiliar] {m}", flush=True)


def main():
    fonte = next((p for p in CANDIDATOS if p and os.path.exists(p)), None)
    if not fonte:
        log("AUXILIAR viva não encontrada — nada feito (o robô segue com a cópia do repo).")
        return 0
    wb = openpyxl.load_workbook(fonte, read_only=True, data_only=True)
    ws = wb["Operacoes_1"] if "Operacoes_1" in wb.sheetnames else wb[wb.sheetnames[0]]
    linhas = list(ws.iter_rows(values_only=True))
    wb.close()
    if not linhas:
        log("AUXILIAR viva vazia — abortado.")
        return 1
    cab = [str(c).strip() if c is not None else "" for c in linhas[0]]
    manter = [i for i, c in enumerate(cab) if c and not _sensivel(c)]
    obrig = {"ufv", "cliente", "responsavel o&m"}
    if not obrig <= {_sem_acento(cab[i]) for i in manter}:
        log(f"cabeçalho inesperado (faltam {obrig}) — abortado, nada gravado.")
        return 1
    dados = [[cab[i] for i in manter]]
    for r in linhas[1:]:
        if not any(v not in (None, "") for v in r):
            continue
        dados.append([(r[i] if i < len(r) else None) for i in manter])
    # só regrava quando muda (evita commit a cada 15 min)
    h = hashlib.sha256(json.dumps(dados, default=str, ensure_ascii=False).encode("utf-8")).hexdigest()
    anterior = open(MARCA).read().strip() if os.path.exists(MARCA) else ""
    if h == anterior and all(os.path.exists(d) for d in DESTINOS):
        log("sem mudanças na AUXILIAR viva.")
        return 0
    for dst in DESTINOS:
        out = openpyxl.Workbook()
        o = out.active
        o.title = "Operacoes_1"
        for row in dados:
            o.append(row)
        out.save(dst)
    open(MARCA, "w").write(h)
    log(f"atualizada: {len(dados) - 1} usinas, {len(manter)} colunas "
        f"({len(cab) - len(manter)} sensíveis removidas) ← {fonte}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
